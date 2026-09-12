"""Recompute convention checks and redraw the figures revised for paper v02."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from numba import njit
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = Path(__file__).resolve().parent / "data"
OUT = Path(__file__).resolve().parent
R0 = 0.02
H, HM = 25.0, 8e-7


def read_xlsx(path: Path) -> np.ndarray:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))[1:]
    wb.close()
    return np.asarray(rows, dtype=float)


def pchip_slopes(x: np.ndarray, y: np.ndarray) -> np.ndarray:
    h = np.diff(x)
    delta = np.diff(y) / h
    d = np.zeros_like(y)
    for i in range(1, len(y) - 1):
        if delta[i - 1] * delta[i] > 0:
            w1, w2 = 2 * h[i] + h[i - 1], h[i] + 2 * h[i - 1]
            d[i] = (w1 + w2) / (w1 / delta[i - 1] + w2 / delta[i])
    d[0] = ((2 * h[0] + h[1]) * delta[0] - h[0] * delta[1]) / (h[0] + h[1])
    d[-1] = ((2 * h[-1] + h[-2]) * delta[-1] - h[-1] * delta[-2]) / (h[-1] + h[-2])
    if d[0] * delta[0] <= 0:
        d[0] = 0
    elif abs(d[0]) > 3 * abs(delta[0]):
        d[0] = 3 * delta[0]
    if d[-1] * delta[-1] <= 0:
        d[-1] = 0
    elif abs(d[-1]) > 3 * abs(delta[-1]):
        d[-1] = 3 * delta[-1]
    return d


@njit(cache=True)
def interp_linear(t, x, y):
    if t <= x[0]:
        return y[0]
    if t >= x[-1]:
        return y[-1]
    j = np.searchsorted(x, t) - 1
    return y[j] + (y[j + 1] - y[j]) * (t - x[j]) / (x[j + 1] - x[j])


@njit(cache=True)
def interp_pchip(t, x, y, d):
    if t <= x[0]:
        return y[0]
    if t >= x[-1]:
        return y[-1]
    j = np.searchsorted(x, t) - 1
    h = x[j + 1] - x[j]
    z = (t - x[j]) / h
    return ((2*z**3 - 3*z**2 + 1) * y[j]
            + (z**3 - 2*z**2 + z) * h * d[j]
            + (-2*z**3 + 3*z**2) * y[j + 1]
            + (z**3 - z**2) * h * d[j + 1])


@njit(cache=True)
def geometry(n):
    x = np.linspace(0.0, 1.0, n + 1)
    faces = R0 * (10.0*x - x**10) / 9.0
    volumes = (faces[1:]**2 - faces[:-1]**2) / 2.0
    centers = (2.0/3.0) * (faces[1:]**3 - faces[:-1]**3) / (faces[1:]**2 - faces[:-1]**2)
    return faces, centers, volumes


@njit(cache=True)
def props(T, C, question, a, k, D):
    for i in range(len(C)):
        if question == 3:
            a[i] = (650 + 128*C[i]) * (1450 + 2736*C[i]/(C[i] + 1))
            k[i] = 0.21 + 0.38*C[i]/(C[i] + 1)
            D[i] = 2.4e-3*np.exp(-0.45/C[i])*np.exp(-3850/T[i])
        else:
            a[i] = (760 + 90*C[i]) * (1850 + 2150*C[i]/(C[i] + 1))
            k[i] = 0.12 + 0.20*C[i]/(C[i] + 1)
            D[i] = 4.2e-4*np.exp(-0.30/C[i])*np.exp(-3850/T[i])


@njit(cache=True)
def links(coef, exchange, scale, faces, centers, out):
    # Reference-domain coefficients: internal b/s^2, surface beta/s.
    for j in range(1, len(coef)):
        out[j] = faces[j] / ((faces[j]-centers[j-1])/(coef[j-1]/scale**2)
                             + (centers[j]-faces[j])/(coef[j]/scale**2))
    out[0] = 0.0
    out[-1] = faces[-1] / ((faces[-1]-centers[-1])/(coef[-1]/scale**2)
                            + 1.0/(exchange/scale))


@njit(cache=True)
def tridiagonal_step(old, capacity, volume, conductance, external, dt, out, upper, rhs):
    n = len(old)
    for i in range(n):
        q = dt/(capacity[i]*volume[i])
        left, right = conductance[i]*q, conductance[i+1]*q
        diagonal = 1 + left + right
        value = 0.0
        if i:
            value += left*(old[i-1]-old[i])
            diagonal -= left*upper[i-1]
            value += left*rhs[i-1]
        if i < n-1:
            value += right*(old[i+1]-old[i])
        else:
            value += right*(external-old[i])
        upper[i] = right/diagonal
        rhs[i] = value/diagonal
    inc = rhs[-1]
    out[-1] = old[-1] + inc
    for i in range(n-2, -1, -1):
        inc = rhs[i] + upper[i]*inc
        out[i] = old[i] + inc


@njit(cache=True)
def center_value(values, faces):
    q0 = 0.5*(faces[0]**2 + faces[1]**2)
    q1 = 0.5*(faces[1]**2 + faces[2]**2)
    return values[0] - (values[1]-values[0])*q0/(q1-q0)


@njit(cache=True)
def simulate(question, criterion_mean, last_env, radius_linear, dt_factor,
             env_t, env_T, env_C, rad_t, rad_R, rad_d, n=120):
    faces, centers, volumes = geometry(n)
    T, C = np.full(n, 301.15), np.full(n, 2.55)
    a, k, D = np.empty(n), np.empty(n), np.empty(n)
    gt, gc = np.empty(n+1), np.empty(n+1)
    nextT, nextC = np.empty(n), np.empty(n)
    guessT, guessC = np.empty(n), np.empty(n)
    upper, rhs = np.empty(n), np.empty(n)
    ones = np.ones(n)
    max_time = 600000.0
    history_t = np.empty(12000)
    history_center = np.empty(12000)
    history_mean = np.empty(12000)
    history_surface = np.empty(12000)
    count, t, next_save = 0, 0.0, 0.0
    old_event = 2.55 - 0.15
    while t < max_time:
        base_dt = 5.0 if t < 600 else (15.0 if t < 3600 else 60.0)
        dt = base_dt*dt_factor
        if t + dt > next_save:
            dt = next_save - t if next_save > t else dt
        dt = min(dt, max_time-t)
        target = t + dt
        if target <= 14400:
            extT = interp_linear(target, env_t, env_T) + 273.15
            extC = interp_linear(target, env_t, env_C)
        elif last_env:
            extT, extC = env_T[-1] + 273.15, env_C[-1]
        else:
            extT, extC = 323.15, 0.05
        if question == 4:
            radius = interp_linear(target, rad_t, rad_R) if radius_linear else interp_pchip(target, rad_t, rad_R, rad_d)
            scale = radius/R0
        else:
            scale = 1.0
        guessT[:], guessC[:] = T, C
        for _ in range(80):
            props(guessT, guessC, question, a, k, D)
            links(k, H, scale, faces, centers, gt)
            tridiagonal_step(T, a, volumes, gt, extT, dt, nextT, upper, rhs)
            props(nextT, guessC, question, a, k, D)
            links(D, HM, scale, faces, centers, gc)
            tridiagonal_step(C, ones, volumes, gc, extC, dt, nextC, upper, rhs)
            change = max(np.max(np.abs(nextT-guessT)), np.max(np.abs(nextC-guessC)))
            guessT[:], guessC[:] = nextT, nextC
            if change < 1e-10:
                break
        T[:], C[:] = nextT, nextC
        t = target
        mean = np.sum(volumes*C)/np.sum(volumes)
        center = center_value(C, faces)
        surface = (C[-1] + (faces[-1]-centers[-1])/(D[-1]/scale**2)*(HM/scale)*extC) / (1 + (faces[-1]-centers[-1])/(D[-1]/scale**2)*(HM/scale))
        event = (mean if criterion_mean else max(center, np.max(C), surface)) - 0.15
        if t >= next_save - 1e-9:
            history_t[count], history_center[count] = t, center
            history_mean[count], history_surface[count] = mean, surface
            count += 1
            next_save += 60.0
        if event < 0:
            finish = t - dt*event/(event-old_event)
            return finish, history_t[:count], history_center[:count], history_mean[:count], history_surface[:count]
        old_event = event
    raise RuntimeError("threshold not reached")


def run_comparisons():
    env = read_xlsx(DATA/"attachment1.xlsx")
    rad = read_xlsx(DATA/"attachment2.xlsx")
    rad_t, rad_R = rad[:, 0], rad[:, 1]/100
    rad_d = pchip_slopes(rad_t, rad_R)
    cases = [
        ("Q3 本文口径", 3, False, False, False),
        ("Q3 附件末值延续", 3, False, True, False),
        ("Q3 平均含水率判据", 3, True, False, False),
        ("Q4 本文口径", 4, False, False, False),
        ("Q4 固定半径", 5, False, False, False),
        ("Q4 附件末值延续", 4, False, True, False),
        ("Q4 平均含水率判据", 4, True, False, False),
        ("Q4 半径线性插值", 4, False, False, True),
    ]
    rows, histories = [], {}
    for label, q, mean, last, linear in cases:
        values = []
        factors = (1.0, 0.5, 0.25, 0.125)
        for factor in factors:
            result = simulate(q, mean, last, linear, factor, env[:, 0], env[:, 1], env[:, 2], rad_t, rad_R, rad_d)
            values.append(result[0]/3600)
            if factor == factors[-1]:
                histories[label] = result
        rows.append({"方案": label, "dt倍率1终点_h": values[0], "dt倍率0.5终点_h": values[1],
                     "dt倍率0.25终点_h": values[2], "dt倍率0.125终点_h": values[3],
                     "末两级时间加密差_h": values[3]-values[2]})
    with (OUT/"口径对比补算.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=rows[0]); w.writeheader(); w.writerows(rows)
    (OUT/"口径对比补算.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return rows, histories


def style():
    plt.rcParams.update({"font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
                         "axes.unicode_minus": False, "font.size": 9, "axes.labelsize": 9,
                         "legend.fontsize": 8, "xtick.labelsize": 8, "ytick.labelsize": 8,
                         "mathtext.fontset": "stix"})


def save(fig, path):
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def redraw(histories):
    style()
    with np.load(DATA/"q1_solution.npz") as z:
        r, T, C = z["radii"]*100, z["temperature_K"]-273.15, z["moisture"]
        times = [100, 300, 600, 900, 1200, 1500, 1800]
        fig, ax = plt.subplots(1, 2, figsize=(8.4, 3.3), constrained_layout=True)
        colors = plt.cm.viridis(np.linspace(.05, .95, len(times)))
        for i, t in enumerate(times):
            ls = "-" if i % 2 == 0 else "--"
            ax[0].plot(r, T[t], color=colors[i], ls=ls, label=f"{t} s")
            ax[1].plot(r, C[t], color=colors[i], ls=ls)
        ax[0].set(xlabel="径向位置 $r$ / cm", ylabel="温度 $T$ / ℃")
        ax[1].set(xlabel="径向位置 $r$ / cm", ylabel="干基含水率 $C$ / (kg/kg)")
        fig.legend(loc="outside upper center", ncol=7, frameon=False)
        save(fig, ROOT/"q1"/"figures"/"fig3_q1_profiles.pdf")
    with np.load(DATA/"q2_solution.npz") as z:
        t, r = z["times"]/3600, z["radii"]*100
        T, C = z["temperature_K"]-273.15, z["moisture"]
        fig, ax = plt.subplots(2, 2, figsize=(8.4, 6.0), constrained_layout=True)
        hours = np.arange(.5, 3.01, .5)
        colors = plt.cm.plasma(np.linspace(.05, .9, len(hours)))
        for h, color in zip(hours, colors): ax[0,0].plot(r, T[int(h*3600)], color=color, label=f"{h:g} h")
        im=ax[0,1].pcolormesh(t, r, C.T, shading="auto", cmap="YlGnBu_r", rasterized=True)
        fig.colorbar(im, ax=ax[0,1], label="$C$ / (kg/kg)")
        for j, label in zip((0,50,100),("中心","1 cm","表面")):
            ax[1,0].plot(t,T[:,j],label=label); ax[1,1].plot(t,C[:,j],label=label)
        ax[0,0].set(xlabel="径向位置 $r$ / cm",ylabel="温度 $T$ / ℃"); ax[0,0].legend(ncol=2)
        ax[0,1].set(xlabel="时间 $t$ / h",ylabel="径向位置 $r$ / cm")
        ax[1,0].set(xlabel="时间 $t$ / h",ylabel="温度 $T$ / ℃"); ax[1,0].legend()
        ax[1,1].set(xlabel="时间 $t$ / h",ylabel="$C$ / (kg/kg)"); ax[1,1].legend()
        for a, tag in zip(ax.flat,"abcd"): a.text(.02,.96,f"({tag})",transform=a.transAxes,va="top",fontweight="bold")
        save(fig, ROOT/"q2"/"figures"/"fig4_q2_evolution.pdf")
    with np.load(DATA/"q3_solution.npz") as z:
        t=z["times"]/3600; C=z["moisture"]; maximum=z["max_C"]; mean=z["mean_C"]
        fig,ax=plt.subplots(1,2,figsize=(8.4,3.3),constrained_layout=True)
        ax[0].plot(t,maximum,label="全域最大值"); ax[0].plot(t,mean,"--",label="径向平均值")
        for j,label in zip((0,50,100),("中心","1 cm","表面")): ax[1].plot(t,C[:,j],label=label)
        for a in ax:
            a.axhline(.15,color="#c44e52",ls=":",label="阈值 0.15"); a.axvline(57.4723,color="#555555",ls="--")
            a.set(xlabel="从题设初态起算的时间 $t$ / h",ylabel="$C$ / (kg/kg)"); a.legend()
        ax[0].text(.02,.96,"(a)",transform=ax[0].transAxes,va="top",fontweight="bold")
        ax[1].text(.02,.96,"(b)",transform=ax[1].transAxes,va="top",fontweight="bold")
        save(fig,ROOT/"q3"/"figures"/"fig5_q3_drying_endpoint.pdf")
    fig,ax=plt.subplots(1,2,figsize=(8.4,3.3),constrained_layout=True)
    for label,color in (("Q4 本文口径","#4c72b0"),("Q4 固定半径","#c44e52")):
        finish,t,center,mean,surface=histories[label]
        ax[0].plot(t/3600,center,label=label,color=color)
    q4=histories["Q4 本文口径"]
    for series,label in ((q4[2],"中心"),(q4[3],"径向平均"),(q4[4],"表面")):
        ax[1].plot(q4[1]/3600,series,label=label)
    for a in ax:
        a.axhline(.15,color="#c44e52",ls=":"); a.set(xlabel="时间 $t$ / h",ylabel="$C$ / (kg/kg)"); a.legend()
    ax[0].text(.02,.96,"(a)",transform=ax[0].transAxes,va="top",fontweight="bold")
    ax[1].text(.02,.96,"(b)",transform=ax[1].transAxes,va="top",fontweight="bold")
    save(fig,ROOT/"q4"/"figures"/"fig7_q4_shrinking_comparison.pdf")
    rows=list(csv.DictReader((ROOT/"analysis"/"sensitivity_results.csv").open(encoding="utf-8-sig")))
    labels=[r["label"] for r in rows if r["level"]=="low"]
    low=np.array([float(r["finish_time_h"]) for r in rows if r["level"]=="low"])
    high=np.array([float(r["finish_time_h"]) for r in rows if r["level"]=="high"])
    fig,ax=plt.subplots(figsize=(7.6,3.5),constrained_layout=True)
    y=np.arange(len(labels)); ax.hlines(y,low,high,color="#bbbbbb"); ax.scatter(low,y,marker="s",label="低水平"); ax.scatter(high,y,label="高水平")
    ax.axvline(float(rows[0]["finish_time_h"]),color="#555555",ls="--",label="基准")
    ax.set(xlabel="达标时间 / h",yticks=y,yticklabels=[s.replace("h_m","$h_m$") for s in labels]); ax.legend(ncol=3)
    save(fig,ROOT/"analysis"/"fig9_sensitivity.pdf")


if __name__ == "__main__":
    comparison_rows, saved_histories = run_comparisons()
    redraw(saved_histories)
    print(json.dumps(comparison_rows, ensure_ascii=False, indent=2))
