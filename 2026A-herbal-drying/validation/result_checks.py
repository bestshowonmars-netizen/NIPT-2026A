"""Independent physical/invariant checks and full Excel-to-array comparison."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bootstrap
import json,math
import numpy as np
from openpyxl import load_workbook
from common.data_loader import load_environment
from common.interpolation import environment_at
from common.material_properties import properties
from common.finite_volume import geometry,integrate_block,reconstruct,links,linear_step
from common.boundary_conditions import surface_value
from common.solver import load_solution
ROOT=bootstrap.ROOT

def core_checks():
    checks={}
    env=load_environment(ROOT/'data/attachment1.xlsx')
    assert env.shape==(241,3) and env[0,0]==0 and env[-1,0]==14400
    for row in env:
        assert np.allclose(environment_at(row[0],env),row[1:],rtol=0,atol=1e-12)
    mid=environment_at(30,env)
    assert np.allclose(mid,(env[0,1:]+env[1,1:])/2,rtol=0,atol=1e-12)
    assert environment_at(14401,env)==(323.15,.05)
    checks['input_interpolation']='241 observations, all knots, midpoint and post-4h convention checked'
    for T,C in [(301.15,2.55),(323.15,.5)]:
        a,k,D=properties(T,C,1)
        assert a==820*2600 and k==.36
        assert math.isclose(D,7e-9*math.exp(-.89/C),rel_tol=1e-14)
        a,k,D=properties(T,C,2)
        expected=((650+128*C)*(1450+2736*C/(C+1)),.21+.38*C/(C+1),.0024*math.exp(-.45/C-3850/T))
        assert np.allclose([a,k,D],expected,rtol=1e-13,atol=0)
    checks['material_formulas']='appendices 2 and 3 checked independently, temperature in K'
    faces,centers,V=geometry(40,.02);r=np.linspace(0,.02,21)
    assert abs(V.sum()-.02**2/2)<1e-18 and np.all(np.diff(faces)>0)
    # The first two exact cell averages of u=2+3r^2 must reconstruct u(0)=2.
    means=2+3*(faces[1:]**2+faces[:-1]**2)/2
    rec=reconstruct(means,np.ones(40),0.,2.,faces,centers,np.array([0.]))
    assert abs(rec[0]-2)<1e-14
    checks['center_reconstruction']='even quadratic volume-average reconstruction checked'
    value=surface_value(2.1,.05,7e-9,8e-7,faces[-1]-centers[-1])
    flux=-7e-9*(value-2.1)/(faces[-1]-centers[-1])
    assert abs(flux-8e-7*(value-.05))<1e-18
    checks['robin_surface']='surface reconstruction satisfies both half-cell and Robin flux'
    eqenv=np.array([[0.,301.15,2.55],[60.,301.15,2.55]])
    for q in [1,2]:
        T,C,ot,oc,log=integrate_block(np.full(40,301.15),np.full(40,2.55),0.,10,.1,q,
            eqenv,faces,centers,V,r,25.,8e-7,1e-10,1e-11,60)
        assert np.max(abs(T-301.15))<1e-12 and np.max(abs(C-2.55))<1e-12
    checks['equilibrium']='both nonlinear models preserve uniform matched-environment states'
    C0=1.2+.7*(centers/.02)**2;T0=np.full(40,310.)
    T,C,ot,oc,log=integrate_block(T0,C0.copy(),0.,10,.05,2,eqenv,faces,centers,V,r,0.,0.,1e-10,1e-11,60)
    mass_error=abs(np.dot(C,V)-np.dot(C0,V))/V.sum()
    assert mass_error<1e-10 and np.max(abs(T-310))<1e-11 and np.max(abs(log[:,1]))==0
    checks['closed_boundary_mass_error']=float(mass_error)
    # Dense independent solve of one frozen finite-volume system checks the Thomas elimination.
    coef=np.linspace(.2,.5,40);a=np.ones(40)*2e6;old=np.linspace(305,310,40)
    g=np.empty(41);links(coef,25.,faces,centers,g)
    dt=.2;A=np.diag(a*V/dt+g[:-1]+g[1:])
    for i in range(39):A[i,i+1]=A[i+1,i]=-g[i+1]
    rhs=a*V/dt*old;rhs[-1]+=g[-1]*320
    expected=np.linalg.solve(A,rhs);actual=np.empty(40)
    linear_step(old,a,V,g,320.,dt,actual,np.empty((2,40)))
    assert np.max(abs(actual-expected))<1e-10
    checks['linear_solver']='increment Thomas solver agrees with independent dense system'
    (ROOT/'validation/core_checks.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
    return checks

def check_result(q,solution,excel_path):
    checks={'question':q}
    end=1800 if q==1 else 10800
    assert len(solution['times'])==end+1
    assert np.all(solution['temperature_K'][0]==301.15) and np.all(solution['moisture'][0]==2.55)
    assert np.isfinite(solution['temperature_K']).all() and np.isfinite(solution['moisture']).all()
    assert np.min(solution['moisture'])>0
    assert solution['metadata']['max_water_balance_residual']<1e-8
    assert solution['metadata']['max_scaled_residual']<=1
    wb=load_workbook(excel_path,read_only=True,data_only=True)
    assert wb.sheetnames==['温度','水分浓度']
    stride=(len(solution['radii'])-1)//20
    for title,key in [('温度','temperature_K'),('水分浓度','moisture')]:
        ws=wb[title]
        assert ws.max_row==end+1 and ws.max_column==22
        rows=ws.iter_rows();header=next(rows)
        assert header[0].value=='时间\\到药材中心的距离'
        assert np.allclose([c.value for c in header[1:]],np.arange(21)/10,atol=1e-14,rtol=0)
        for t,row in enumerate(rows,1):
            assert row[0].value==t
            expected=solution[key][t,::stride]-(273.15 if key=='temperature_K' else 0)
            for j,cell in enumerate(row[1:]):
                assert isinstance(cell.value,(int,float)) and cell.number_format=='0.0000'
                assert abs(cell.value-round(float(expected[j]),4))<1e-12
    wb.close()
    import csv
    times=[100,300,600,900,1200,1500,1800] if q==1 else [1800,3600,5400,7200,9000,10800]
    for no,key in [(1 if q==1 else 3,'temperature_K'),(2 if q==1 else 4,'moisture')]:
        with (ROOT/f'q{q}/tables/table{no}.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))[1:]
        assert len(rows)==len(times)
        for t,row in zip(times,rows):
            assert float(row[0])==(t if q==1 else t/3600)
            expected=solution[key][t,::((len(solution['radii'])-1)//4)]-(273.15 if key=='temperature_K' else 0)
            assert row[1:]==[f'{x:.4f}' for x in expected]
    checks.update({'all_excel_values_checked':end*21*2,'water_balance':solution['metadata']['max_water_balance_residual'],
        'minimum_C':float(np.min(solution['moisture'])),'tables_match':True,'initial_state_independent':True})
    return checks

if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--core',action='store_true');args=p.parse_args()
    report={'core':core_checks()}
    if not args.core:
        report['results']=[check_result(q,load_solution(ROOT/f'q{q}/solution.npz'),ROOT/f'submission_results/result{q}.xlsx') for q in [1,2]]
    (ROOT/'validation/result_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
