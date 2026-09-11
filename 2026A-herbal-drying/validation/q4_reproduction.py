"""Check a complete standalone Q4 solve against unrounded formal outputs."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bootstrap
import argparse
import hashlib
import json
import subprocess
import numpy as np
from openpyxl import load_workbook
from common.solver import load_solution

ROOT=bootstrap.ROOT


def verify():
    reference=load_solution(ROOT/'q4/solution.npz')
    folder=ROOT/'paper_appendix_minimal/reproduced/q4'
    reproduced=load_solution(folder/'q4_independent.npz')
    assert reproduced['metadata']['from_initial']
    assert reference['metadata']['finish_time_s']==reproduced['metadata']['finish_time_s']
    assert reference['metadata']['bracket']==reproduced['metadata']['bracket']
    assert np.array_equal(reference['times'],reproduced['times'])
    assert np.array_equal(reference['valid_mask'][:,::5],reproduced['domain_mask'])
    differences={}
    for key in ('final_T','final_C','radius_m','surface_temperature_K','surface_moisture',
                'max_C','max_r','water_balance','loss_cumulative'):
        difference=float(np.max(abs(reference[key]-reproduced[key])))
        assert difference<=1e-12,(key,difference)
        differences[key]=difference
    for key in ('temperature_K','moisture'):
        original=reference[key][:,::5]
        candidate=reproduced[key]
        assert np.array_equal(np.isfinite(original),np.isfinite(candidate))
        difference=float(np.max(abs(original[np.isfinite(original)]-candidate[np.isfinite(candidate)])))
        assert difference<=1e-12,(key,difference)
        differences[key]=difference
    original_book=load_workbook(ROOT/'submission_results/result4.xlsx',data_only=True,read_only=True)
    new_book=load_workbook(folder/'result4.xlsx',data_only=True,read_only=True)
    original_rows=list(original_book.active.values);new_rows=list(new_book.active.values)
    original_book.close();new_book.close()
    assert original_rows==new_rows,'Standalone Excel differs in an output cell or exterior blank'
    report={'passed':True,'same_algorithm_from_original_initial_state':True,
        'independent_model_validation':False,'finish_time_s':reference['metadata']['finish_time_s'],
        'maximum_differences':differences,'excel_rows_including_header':len(original_rows),
        'excel_values_and_blanks_identical':True,'source_hashes':{}}
    for relative in ('paper_appendix_minimal/minimal_solver.py','paper_appendix_minimal/minimal_q4.py',
                     'paper_appendix_minimal/reproduced/q4/q4_independent.npz',
                     'paper_appendix_minimal/reproduced/q4/result4.xlsx',
                     'data/attachment1.xlsx','data/attachment2.xlsx',
                     'q4/solution.npz','submission_results/result4.xlsx'):
        report['source_hashes'][relative]=hashlib.sha256((ROOT/relative).read_bytes()).hexdigest()
    (ROOT/'validation/q4_reproduction.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--run',action='store_true')
    args=parser.parse_args()
    if args.run:
        from validation.build_q4_appendix import build
        target=build()
        subprocess.run([sys.executable,'-B',str(target)],cwd=ROOT,check=True)
    verify()
