"""Q3 invariants, restart equivalence and full deliverable-to-array checks.

Refinement and serialization checks are numerical verification, not comparison
with a true solution or measured internal moisture/temperature fields.
"""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import bootstrap
import argparse
import csv
import json
import math
import tempfile
import hashlib
import numpy as np
from openpyxl import load_workbook
from common.solver import load_solution
from common.data_loader import load_environment
from validation.evidence_paths import project_path

ROOT = bootstrap.ROOT
LIMITS = {"field_C": 2e-5, "endpoint_s": .1, "event_s": .02,
          "restart": 1e-12, "water": 1e-8}


def matching_indices(values, requested, tolerance=1e-8):
    """Match existing coordinates only; never interpolate validation data."""
    values, requested = np.asarray(values), np.asarray(requested)
    indices = np.searchsorted(values, requested)
    indices = np.minimum(indices, len(values)-1)
    previous = np.maximum(indices-1, 0)
    choose_previous = abs(values[previous]-requested) < abs(values[indices]-requested)
    indices[choose_previous] = previous[choose_previous]
    if not np.all(abs(values[indices]-requested) <= tolerance):
        raise AssertionError("Required comparison coordinates are absent; no interpolation allowed")
    return indices


def independent_maximum(solution, environment):
    """Re-evaluate all cell values plus independently reconstructed endpoints."""
    C, T = solution["final_C"], solution["final_T"]
    faces, centers = solution["faces"], solution["centers"]
    q0, q1 = (faces[:2]**2+faces[1:3]**2)/2
    center = C[0]-(C[1]-C[0])*q0/(q1-q0)
    t = float(solution["times"][-1])
    ext = float(np.interp(t, environment[:, 0], environment[:, 2])) if t <= environment[-1, 0] else .05
    hm = float(solution["metadata"].get("hm", 8e-7))
    diffusivity = 2.4e-3*math.exp(-.45/float(C[-1])-3850/float(T[-1]))
    delta = faces[-1]-centers[-1]
    surface = C[-1]+hm*delta/(diffusivity+hm*delta)*(ext-C[-1])
    values = np.r_[center, C, surface]
    positions = np.r_[0., centers, faces[-1]]
    i = int(np.argmax(values))
    return float(values[i]), float(positions[i])


def check_solution(solution, source, require_dense=False):
    meta = solution["metadata"]
    assert meta["question"] == 3 and meta["model_question"] == 2
    assert meta["status"] == "complete" and meta["threshold"] == .15
    source_path = project_path(meta["source_path"], ROOT).resolve()
    source_digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
    recorded_digests = [expected for recorded, expected in meta["source_hashes"].items()
                        if project_path(recorded, ROOT).resolve() == source_path]
    assert recorded_digests and all(value == source_digest for value in recorded_digests), \
        "The relocated Q2 source must match its original recorded hash"
    times, radii = solution["times"], solution["radii"]
    finish = float(meta["finish_time_s"])
    tol = float(meta.get("epsilon_t", meta.get("event_tol", .01)))
    assert len(times) > 181 and times[0] == 0 and np.all(np.diff(times) > 0)
    assert abs(times[-1]-finish) < 1e-9
    regular = np.arange(0., math.floor((finish+1e-9)/60)*60+1, 60.)
    expected = regular if abs(regular[-1]-finish) < 1e-8 else np.r_[regular, finish]
    assert np.array_equal(times, expected), "Output must contain every 60 s plus exact endpoint"
    required_r = np.linspace(0., .02, 21)
    cols = matching_indices(radii, required_r, 1e-12)
    if require_dense:
        assert len(radii) == 101 and np.allclose(radii, np.linspace(0., .02, 101), rtol=0, atol=1e-14)
    for key in ("temperature_K", "moisture"):
        assert solution[key].shape == (len(times), len(radii))
        assert np.isfinite(solution[key]).all() and np.min(solution[key]) > 0
    assert np.all(solution["temperature_K"][0] == 301.15)
    assert np.all(solution["moisture"][0] == 2.55)
    assert np.max(solution["moisture"]) <= 2.55+1e-10
    for key in ("faces", "centers", "volumes"):
        assert np.array_equal(solution[key], source[key]), "Q2 mesh must be continued without projection"
    assert len(solution["final_C"]) == int(meta["cells"])
    for key in ("final_T", "final_C"):
        assert np.isfinite(solution[key]).all() and np.min(solution[key]) > 0
    assert source["metadata"]["question"] == 2
    assert float(source["times"][-1]) == 10800 and float(meta["start_time_s"]) == 10800
    prefix = times <= 10800
    source_rows = matching_indices(source["times"], times[prefix])
    source_cols = matching_indices(source["radii"], required_r, 1e-12)
    prefix_difference = {}
    for key in ("temperature_K", "moisture"):
        error = float(np.max(abs(solution[key][prefix][:, cols]-source[key][source_rows][:, source_cols])))
        assert error <= 1e-12, "Q2 output prefix changed"
        prefix_difference[key] = error
    logs = solution["logs"]
    assert logs.shape == (len(times), 9) and np.isfinite(logs).all()
    assert np.max(logs[:, 7]) <= 1 and np.min(logs[:, 1]) >= -1e-12
    cumulative = np.cumsum(logs[:, 1])
    assert np.max(abs(cumulative-solution["loss_cumulative"])) <= 1e-10
    row3h = int(matching_indices(times, np.array([10800.]))[0])
    prefix_loss_error = abs(float(solution["loss_cumulative"][row3h])-float(source["logs"][:, 1].sum()))
    assert prefix_loss_error <= 1e-10, "Q3 cumulative loss must include the whole Q2 history"
    expected_balance = solution["mean_C"]+solution["loss_cumulative"]-2.55
    assert np.max(abs(expected_balance-solution["water_balance"])) <= 1e-12
    water_error = float(np.max(abs(expected_balance)))
    assert water_error <= LIMITS["water"]
    mean_final = float(np.dot(solution["final_C"], solution["volumes"])/solution["volumes"].sum())
    assert abs(mean_final-solution["mean_C"][-1]) <= 1e-12
    assert np.max(abs(logs[1:, 0]-solution["mean_C"][1:])) <= 1e-12
    assert np.isnan(solution["max_C"][times < 10800]).all()
    assert np.isnan(solution["max_r"][times < 10800]).all()
    post = times >= 10800
    assert np.isfinite(solution["max_C"][post]).all()
    assert np.all(solution["max_C"][post] >= np.max(solution["moisture"][post], axis=1)-1e-12)
    env = load_environment(ROOT/"data/attachment1.xlsx")
    independent_M, independent_r = independent_maximum(solution, env)
    assert independent_M < .15, "The unrounded full-grid maximum must be strictly below 0.15"
    assert abs(independent_M-float(solution["max_C"][-1])) <= 1e-12
    assert abs(independent_r-float(solution["max_r"][-1])) <= 1e-12
    bracket = meta["bracket"]
    assert bracket["g_minus"] >= 0 and bracket["g_plus"] < 0
    assert 0 < bracket["width"] <= tol+1e-10
    assert abs(bracket["width"]-(bracket["t_plus"]-bracket["t_minus"])) <= 1e-10
    assert abs(bracket["t_plus"]-finish) <= 1e-9
    assert abs(bracket["g_plus"]-(independent_M-.15)) <= 1e-12
    from common.continuation import advance_interval
    anchor_time = float(solution["event_anchor_time_s"])
    anchor_loss = float(solution["event_anchor_loss_cumulative"])
    assert anchor_time <= bracket["t_minus"] < finish
    endpoint_state_checks = {}
    for side, Tkey, Ckey in (("minus", "event_lower_T", "event_lower_C"),
                            ("plus", "final_T", "final_C")):
        endpoint_time = bracket[f"t_{side}"]
        check_state = {**solution, "final_T": solution[Tkey], "final_C": solution[Ckey],
                       "times": np.array([endpoint_time])}
        endpoint_maximum, _ = independent_maximum(check_state, env)
        assert abs(endpoint_maximum-.15-bracket[f"g_{side}"]) <= 1e-12
        assert endpoint_maximum >= .15 if side == "minus" else endpoint_maximum < .15
        newT, newC, diag = advance_interval(solution["event_anchor_T"], solution["event_anchor_C"],
            anchor_time, endpoint_time, float(meta["event_dt"]), env, solution["faces"],
            solution["centers"], solution["volumes"], float(meta["h"]), float(meta["hm"]),
            float(meta["atol"]), float(meta["rtol"]), int(meta["max_iterations"]))
        error = max(float(np.max(abs(newT-solution[Tkey]))), float(np.max(abs(newC-solution[Ckey]))))
        assert error <= LIMITS["restart"], "Endpoint state is not the claimed integrated anchor trajectory"
        stored_loss = float(solution["event_lower_loss_cumulative"]) if side == "minus" else float(solution["loss_cumulative"][-1])
        loss_error = abs(anchor_loss+float(diag[1])-stored_loss)
        assert loss_error <= LIMITS["restart"]
        endpoint_state_checks[side] = {"independent_max_C": endpoint_maximum,
                                      "reintegrated_state_max_difference": error,
                                      "reintegrated_loss_difference": loss_error}
    trace = solution["event_trace"]
    assert trace.ndim == 2 and trace.shape[1] == 4 and np.isfinite(trace).all()
    assert np.max(abs(trace[:, 1]-.15-trace[:, 3])) <= 1e-14
    for side in ("minus", "plus"):
        matches = abs(trace[:, 0]-bracket[f"t_{side}"]) < 1e-9
        assert np.any(matches & (abs(trace[:, 3]-bracket[f"g_{side}"]) < 1e-12))
    return {"passed": True, "finish_time_s": finish, "strict_final_max_C": independent_M,
            "strict_final_max_r_m": independent_r, "event_bracket": bracket,
            "endpoint_full_state_checks": endpoint_state_checks,
            "maximum_water_balance_residual": water_error,
            "q2_prefix_max_difference": prefix_difference,
            "q2_prefix_loss_difference": prefix_loss_error,
            "original_mesh_preserved": True, "sampled_C_values_checked": int(len(times)*21),
            "q2_source_sha256_verified": source_digest,
            "interpretation": "Discrete invariants and unchanged supplied-data prefix; not true-solution or experimental validation"}


def check_deliverables(solution, excel_path, table_path, endpoint_path=None):
    """Check every delivered number against the full-precision solution."""
    times, radii = solution["times"], solution["radii"]
    finish = float(solution["metadata"]["finish_time_s"])
    excel_times = np.arange(60., math.floor((finish+1e-9)/60)*60+1, 60.)
    ridx = matching_indices(radii, np.linspace(0., .02, 21), 1e-12)
    tidx = matching_indices(times, excel_times)
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        assert wb.sheetnames == ["Sheet1"]
        ws = wb["Sheet1"]
        assert ws.max_row == len(excel_times)+1 and ws.max_column == 22
        rows = ws.iter_rows()
        header = next(rows)
        assert np.allclose([c.value for c in header[1:]], np.arange(21)/10, rtol=0, atol=1e-12)
        for t, i, row in zip(excel_times, tidx, rows):
            assert row[0].value == t
            for cell, value in zip(row[1:], solution["moisture"][i, ridx]):
                assert isinstance(cell.value, (int, float)) and cell.number_format == "0.0000"
                assert abs(cell.value-round(float(value), 4)) <= 1e-12
    finally:
        wb.close()
    with Path(table_path).open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    table_times = np.arange(21600., math.floor((finish+1e-9)/21600)*21600+1, 21600.)
    if not len(table_times) or abs(table_times[-1]-finish) > 1e-8:
        table_times = np.r_[table_times, finish]
    assert len(rows) == len(table_times)
    table_cols = matching_indices(radii, np.linspace(0., .02, 5), 1e-12)
    value_columns = [f"r={r:g} cm" for r in (0, .5, 1, 1.5, 2)]
    for j, (row, t) in enumerate(zip(rows, table_times)):
        assert abs(float(row["time_s"])-t) <= 1e-9
        assert abs(float(row["time_h"])-t/3600) <= 1e-12
        is_endpoint = j == len(rows)-1
        assert row["row_kind"] == ("endpoint" if is_endpoint else "regular_6h")
        i = int(matching_indices(times, np.array([t]))[0])
        assert [row[k] for k in value_columns] == [f"{v:.4f}" for v in solution["moisture"][i, table_cols]]
    if endpoint_path is not None:
        endpoint = json.loads(Path(endpoint_path).read_text(encoding="utf-8"))
        assert abs(float(endpoint["finish_time_s"])-finish) <= 1e-9
        assert abs(float(endpoint["finish_time_h"])-finish/3600) <= 1e-12
        assert abs(float(endpoint["max_C"])-float(solution["max_C"][-1])) <= 1e-14
        assert abs(float(endpoint["max_r"])-float(solution["max_r"][-1])) <= 1e-14
        assert endpoint["max_C"] < .15
        assert abs(endpoint["g_plus"]-(endpoint["max_C"]-.15)) <= 1e-14
        assert endpoint["bracket"] == solution["metadata"]["bracket"]
        assert np.array_equal(np.asarray(endpoint["radii_m"]), radii)
        assert np.array_equal(np.asarray(endpoint["moisture"]), solution["moisture"][-1])
        for key in ("mean_C", "loss_cumulative", "water_balance"):
            assert abs(endpoint[key]-float(solution[key][-1])) <= 1e-14
    return {"passed": True, "all_excel_C_values_checked": int(len(excel_times)*21),
            "table5_rows_checked": len(rows), "excel_contains_only_regular_minutes": True,
            "endpoint_checked_without_rounding": endpoint_path is not None,
            "interpretation": "Artifact-to-raw-array consistency, not evidence of physical accuracy"}


def core_restart_checks(source_path=None, dt=.25, duration=120., split=60., probe_end=14400.):
    """Compare compatible step schedules, disk restart and absolute-time driving."""
    from common.continuation import advance_interval, global_maximum
    from common.finite_volume import integrate_block
    source_path = project_path(source_path or ROOT/"q2/solution.npz", ROOT)
    source = load_solution(source_path)
    cfg = source["metadata"]
    env = load_environment(ROOT/"data/attachment1.xlsx")
    T0, C0 = source["final_T"].copy(), source["final_C"].copy()
    start = float(source["times"][-1])
    assert start == 10800 and 0 < split < duration
    assert abs(split/dt-round(split/dt)) < 1e-9 and abs(duration/dt-round(duration/dt)) < 1e-9
    args = (env, source["faces"], source["centers"], source["volumes"],
            cfg["h"], cfg["hm"], cfg["atol"], cfg["rtol"], cfg["max_iterations"])
    Ta, Ca, da = advance_interval(T0, C0, start, start+duration, dt, *args)
    Ts, Cs, ds = advance_interval(T0, C0, start, start+split, dt, *args)
    with tempfile.TemporaryDirectory(prefix="q3_restart_", dir=bootstrap.RUNTIME/"tmp") as tmp:
        path = Path(tmp)/"state.npz"
        np.savez_compressed(path, T=Ts, C=Cs, time=start+split, loss=source["logs"][:, 1].sum()+ds[1])
        with np.load(path, allow_pickle=False) as saved:
            Tr, Cr, dr = advance_interval(saved["T"], saved["C"], float(saved["time"]), start+duration, dt, *args)
            restored_prefix_loss = float(saved["loss"])
    restart = {"T": float(np.max(abs(Ta-Tr))), "C": float(np.max(abs(Ca-Cr))),
               "loss": float(abs(da[1]-(ds[1]+dr[1])))}
    assert max(restart.values()) <= LIMITS["restart"]
    assert np.array_equal(T0, source["final_T"]) and np.array_equal(C0, source["final_C"])
    oldT, oldC, _, _, oldlog = integrate_block(T0.copy(), C0.copy(), start, int(start+duration), dt, 2,
        env, source["faces"], source["centers"], source["volumes"], source["radii"],
        cfg["h"], cfg["hm"], cfg["atol"], cfg["rtol"], cfg["max_iterations"])
    kernel = {"T": float(np.max(abs(Ta-oldT))), "C": float(np.max(abs(Ca-oldC))),
              "loss": float(abs(da[1]-oldlog[:, 1].sum()))}
    assert max(kernel.values()) <= LIMITS["restart"]
    Tp, Cp, dp = advance_interval(T0, C0, start, probe_end, dt, *args)
    oldT, oldC, _, _, oldlog = integrate_block(T0.copy(), C0.copy(), start, int(probe_end), dt, 2,
        env, source["faces"], source["centers"], source["volumes"], source["radii"],
        cfg["h"], cfg["hm"], cfg["atol"], cfg["rtol"], cfg["max_iterations"])
    absolute = {"T": float(np.max(abs(Tp-oldT))), "C": float(np.max(abs(Cp-oldC))),
                "loss": float(abs(dp[1]-oldlog[:, 1].sum()))}
    assert max(absolute.values()) <= LIMITS["restart"]
    cumulative_loss = float(source["logs"][:, 1].sum()+dp[1])
    water = float(abs(np.dot(Cp, source["volumes"])/source["volumes"].sum()+cumulative_loss-2.55))
    assert water <= LIMITS["water"]
    # A synthetic interior maximum verifies that termination does not use the center alone.
    fakeC = np.full_like(C0, .12); fakeC[len(fakeC)//2] = .2
    maximum = global_maximum(T0, fakeC, start, env, source["faces"], source["centers"], cfg["h"], cfg["hm"])
    assert maximum[0] == .2 and maximum[1] == source["centers"][len(fakeC)//2]
    return {"passed": True, "source": str(source_path), "step_s": dt,
            "uninterrupted_vs_save_reload": restart, "old_vs_new_kernel": kernel,
            "absolute_time_probe": {"from_s": start, "to_s": probe_end, **absolute},
            "water_balance_including_q2": water, "inputs_unchanged": True,
            "interior_maximum_detected": True, "restored_prefix_loss": restored_prefix_loss,
            "interpretation": "Same equations and same step schedule; serialization/kernel equivalence, not independent truth"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--core", action="store_true")
    parser.add_argument("--solution", type=Path, default=ROOT/"q3/solution.npz")
    parser.add_argument("--source", type=Path, default=ROOT/"q2/solution.npz")
    parser.add_argument("--skip-deliverables", action="store_true")
    parser.add_argument("--output", type=Path, default=ROOT/"validation/q3_checks.json")
    args = parser.parse_args()
    if args.core:
        report = {"core": core_restart_checks(args.source)}
    else:
        solution = load_solution(project_path(args.solution, ROOT))
        report = {"solution": check_solution(solution, load_solution(project_path(args.source, ROOT)), require_dense=True)}
        if not args.skip_deliverables:
            report["deliverables"] = check_deliverables(solution, ROOT/"submission_results/result3.xlsx",
                ROOT/"q3/tables/table5.csv", ROOT/"q3/endpoint.json")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
