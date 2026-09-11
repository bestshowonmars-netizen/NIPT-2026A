"""Q4 checks for moving physical masks, material fields, and dry-mass balance.

Numerical invariants and refinement differences are not measured-field
validation or error bounds against an unknown exact solution.
"""
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import bootstrap
from validation.evidence_paths import project_path as resolve_evidence_path
import argparse
import csv
import hashlib
import json
import math
import tempfile
import numpy as np
from openpyxl import load_workbook
from common.solver import load_solution

ROOT = bootstrap.ROOT
LIMITS = {"field_C": 2e-5, "endpoint_s": .1, "event_s": .02,
          "restart": 1e-12, "water": 1e-8}
RADIUS_0 = .02
INITIAL_T, INITIAL_C = 301.15, 2.55
DOMAIN_ATOL = 64*np.finfo(float).eps*RADIUS_0


def project_path(value):
    """Resolve portable keys and older Windows separator records."""
    return resolve_evidence_path(value, ROOT)


def matching_indices(values, requested, tolerance=1e-9):
    """Select existing coordinates; no validation-time interpolation."""
    values, requested = np.asarray(values), np.asarray(requested)
    assert values.ndim == 1 and len(values) and np.all(np.diff(values) > 0)
    indices = np.minimum(np.searchsorted(values, requested), len(values)-1)
    previous = np.maximum(indices-1, 0)
    use_previous = abs(values[previous]-requested) < abs(values[indices]-requested)
    indices[use_previous] = previous[use_previous]
    assert np.all(abs(values[indices]-requested) <= tolerance), "Missing exact comparison coordinates"
    return indices


def expected_mask(radii, radius_history):
    """Use the unrounded radius and a strict coordinate comparison."""
    return np.asarray(radii)[None, :] <= np.asarray(radius_history)[:, None]


def independent_properties(T, C):
    """Appendix 4 formulas, deliberately independent of the Q4 kernel."""
    T, C = np.asarray(T), np.asarray(C)
    assert np.all(T > 0) and np.all(C > 0)
    fraction = C/(C+1)
    return ((760+90*C)*(1850+2150*fraction), .12+.20*fraction,
            4.2e-4*np.exp(-.30/C-3850/T))


def independent_boundary_values(T, C, faces, centers, radius, extT, extC, h, hm):
    """Reconstruct center and true surface using physical distances."""
    scale = radius/faces[-1]
    q0, q1 = (faces[:2]**2+faces[1:3]**2)/2
    centerT = T[0]-(T[1]-T[0])*q0/(q1-q0)
    centerC = C[0]-(C[1]-C[0])*q0/(q1-q0)
    _, k, D = independent_properties(T[-1:], C[-1:])
    delta = scale*(faces[-1]-centers[-1])
    surfaceT = T[-1]+h*delta/(k[0]+h*delta)*(extT-T[-1])
    surfaceC = C[-1]+hm*delta/(D[0]+hm*delta)*(extC-C[-1])
    candidates = np.r_[centerC, C, surfaceC]
    positions = np.r_[0., scale*centers, radius]
    where = int(np.argmax(candidates))
    return {"center_T": float(centerT), "center_C": float(centerC),
            "surface_T": float(surfaceT), "surface_C": float(surfaceC),
            "max_C": float(candidates[where]), "max_r": float(positions[where])}


def fixed_appendix4_reference(T0, C0, faces, centers, volumes, dt, steps,
                              extT, extC, h, hm, atol=1e-10, rtol=1e-11):
    """Small dense fixed-radius reference for the s=1 degeneration check.

    This independently assembles the same finite-volume equations; it does not
    call Q4 properties, conductances, TDMA, or time stepping. It checks equation
    implementation, not independence of the assumed physical model.
    """
    T, C = np.array(T0, copy=True), np.array(C0, copy=True)
    size = len(T)
    loss, iterations = 0., 0

    def conductances(coefficient, exchange):
        g = np.zeros(size+1)
        resistance = ((faces[1:-1]-centers[:-1])/coefficient[:-1]
                      +(centers[1:]-faces[1:-1])/coefficient[1:])
        g[1:-1] = faces[1:-1]/resistance
        if exchange > 0:
            g[-1] = faces[-1]/((faces[-1]-centers[-1])/coefficient[-1]+1/exchange)
        return g

    def dense_increment(old, capacity, g, external):
        mass = capacity*volumes/dt
        matrix = np.diag(mass+g[:-1]+g[1:])
        matrix += np.diag(-g[1:-1], 1)+np.diag(-g[1:-1], -1)
        rhs = np.zeros(size)
        rhs[:-1] += g[1:-1]*(old[1:]-old[:-1])
        rhs[1:] += g[1:-1]*(old[:-1]-old[1:])
        rhs[-1] += g[-1]*(external-old[-1])
        return old+np.linalg.solve(matrix, rhs)

    def scaled_residual(old, new, capacity, g, external):
        flux = np.zeros(size)
        flux[:-1] += g[1:-1]*(new[1:]-new[:-1])
        flux[1:] += g[1:-1]*(new[:-1]-new[1:])
        flux[-1] += g[-1]*(external-new[-1])
        mass = capacity*volumes/dt
        denominator = (mass+g[:-1]+g[1:])*(atol+rtol*abs(new))
        return np.max(abs(mass*(new-old)-flux)/denominator)

    for _ in range(steps):
        guessT, guessC = T.copy(), C.copy()
        for iteration in range(100):
            a, k, _ = independent_properties(guessT, guessC)
            newT = dense_increment(T, a, conductances(k, h), extT)
            _, _, D = independent_properties(newT, guessC)
            newC = dense_increment(C, np.ones(size), conductances(D, hm), extC)
            change = max(np.max(abs(newT-guessT)/(atol+rtol*abs(newT))),
                         np.max(abs(newC-guessC)/(atol+rtol*abs(newC))))
            if change <= 1:
                a, k, D = independent_properties(newT, newC)
                gt, gc = conductances(k, h), conductances(D, hm)
                residual = max(scaled_residual(T, newT, a, gt, extT),
                               scaled_residual(C, newC, np.ones(size), gc, extC))
                if residual <= 1:
                    break
            guessT, guessC = newT, newC
        else:
            raise AssertionError("Independent fixed-radius reference did not converge")
        iterations += iteration+1
        loss += dt*gc[-1]*(newC[-1]-extC)/volumes.sum()
        T, C = newT, newC
    return T, C, {"loss": float(loss), "steps": steps, "iterations": iterations}


def check_field_domains(solution):
    """Require exact masks and finite data everywhere that really exists."""
    times, radii = solution["times"], solution["radii"]
    radius = solution["radius_m"]
    xi = solution["xi"]
    assert np.all(np.diff(times) > 0) and times[0] == 0
    assert radius.shape == times.shape and np.isfinite(radius).all()
    assert np.all(radius > 0) and np.max(radius) <= RADIUS_0+DOMAIN_ATOL
    assert np.all(np.diff(radius) <= DOMAIN_ATOL)
    mask = solution["valid_mask"]
    assert mask.dtype == np.bool_ and mask.shape == (len(times), len(radii))
    expected = expected_mask(radii, radius)
    assert np.array_equal(mask, expected), "Mask must use unrounded R(t), not displayed radius"
    for key in ("temperature_K", "moisture"):
        values = solution[key]
        assert values.shape == mask.shape
        assert np.isfinite(values[mask]).all() and np.all(values[mask] > 0)
        assert np.isnan(values[~mask]).all(), "Outside material must be NaN, not extrapolated"
    assert np.all(np.diff(xi) > 0) and xi[0] == 0 and xi[-1] == 1
    for key in ("temperature_ref_K", "moisture_ref"):
        assert solution[key].shape == (len(times), len(xi))
        assert np.isfinite(solution[key]).all() and np.min(solution[key]) > 0
    for key in ("surface_temperature_K", "surface_moisture", "max_C", "max_r"):
        assert solution[key].shape == times.shape and np.isfinite(solution[key]).all()
    assert np.max(abs(solution["temperature_ref_K"][:, -1]-solution["surface_temperature_K"])) <= 1e-12
    assert np.max(abs(solution["moisture_ref"][:, -1]-solution["surface_moisture"])) <= 1e-12
    assert np.all(solution["max_C"] >= np.max(solution["moisture_ref"], axis=1)-1e-12)
    assert np.all(solution["max_r"] >= -DOMAIN_ATOL)
    assert np.all(solution["max_r"] <= radius+DOMAIN_ATOL)
    assert np.all(solution["temperature_ref_K"][0] == INITIAL_T)
    assert np.all(solution["moisture_ref"][0] == INITIAL_C)
    assert np.all(solution["temperature_K"][0, mask[0]] == INITIAL_T)
    assert np.all(solution["moisture"][0, mask[0]] == INITIAL_C)
    assert np.max(solution["moisture_ref"]) <= INITIAL_C+1e-10
    # Boundary equality is an actual surface value, not nearest inside sample.
    same_as_surface = radii[None, :] == radius[:, None]
    row, col = np.where(same_as_surface)
    assert np.all(mask[row, col])
    assert np.max(abs(solution["moisture"][row, col]-solution["surface_moisture"][row]), initial=0) <= 1e-12
    assert not expected_mask(np.array([.012]), np.array([.01198]))[0, 0]
    return {"passed": True, "physical_in_domain_values": int(mask.sum()),
            "physical_outside_values": int((~mask).sum()),
            "reference_values_per_field": int(solution["moisture_ref"].size),
            "radius_1_198_cm_excludes_1_2_cm": True,
            "surface_is_separate_and_matches_xi_1": True}


def core_checks(radius_path=None):
    """Small deterministic checks, without a production-length solve."""
    from common.finite_volume import geometry
    from common.data_loader import load_environment
    from q4.kernel import advance_interval, global_maximum, sample_fields, surface_fields
    from q4.shrinking_radius import load_radius_data, radius_at
    radius_path = Path(radius_path or ROOT/"data/attachment2.xlsx")
    data, slopes = load_radius_data(radius_path)
    workbook = load_workbook(radius_path, read_only=True, data_only=True)
    try:
        raw = np.asarray(list(workbook.worksheets[0].values)[1:], dtype=float)
    finally:
        workbook.close()
    assert np.array_equal(data[:, 0], raw[:, 0])
    assert np.max(abs(data[:, 1]-raw[:, 1]/100)) <= 1e-16
    at_nodes = np.asarray([radius_at(float(t), data, slopes) for t in data[:, 0]])
    assert np.max(abs(at_nodes-data[:, 1])) <= 1e-15
    worst_overshoot = 0.
    plateau_intervals = 0
    for i in range(len(data)-1):
        sample_t = np.linspace(data[i, 0], data[i+1, 0], 21)
        sample_R = np.asarray([radius_at(float(t), data, slopes) for t in sample_t])
        worst_overshoot = max(worst_overshoot, float(np.max(sample_R)-data[i, 1]),
                             float(data[i+1, 1]-np.min(sample_R)))
        assert np.all(np.diff(sample_R) <= DOMAIN_ATOL)
        if data[i, 1] == data[i+1, 1]:
            plateau_intervals += 1
            assert np.all(sample_R == data[i, 1]), "A flat radius segment must stay exact"
    assert worst_overshoot <= DOMAIN_ATOL
    assert abs(radius_at(2*data[-1, 0], data, slopes)-data[-1, 1]) <= DOMAIN_ATOL
    assert abs(data[-1, 1]-.01198) <= DOMAIN_ATOL
    assert radius_at(2*data[-1, 0], data, slopes, False) == RADIUS_0
    n = 24
    faces, centers, volumes = geometry(n, RADIUS_0)
    T0 = 310-2*(centers/RADIUS_0)**2
    C0 = 1.2-.2*(centers/RADIUS_0)**2
    env = np.array([[0., 330., .1], [100., 330., .1]])
    atol, rtol, h, hm = 1e-10, 1e-11, 25., 8e-7
    dt, steps, start = .125, 8, 10.
    end = start+dt*steps
    static_comparisons = {}
    for radius in (RADIUS_0, .012):
        fixed_data = np.array([[0., radius], [100., radius]])
        fixed_slopes = np.zeros(2)
        T, C, diag = advance_interval(T0, C0, start, end, dt, env,
            faces, centers, volumes, h, hm, atol, rtol, 60,
            fixed_data, fixed_slopes, True, 1., 0.)
        scale = radius/RADIUS_0
        referenceT, referenceC, reference_diag = fixed_appendix4_reference(
            T0, C0, faces*scale, centers*scale, volumes*scale**2,
            dt, steps, env[0, 1], env[0, 2], h, hm, atol, rtol)
        errorT, errorC = float(np.max(abs(T-referenceT))), float(np.max(abs(C-referenceC)))
        error_loss = abs(float(diag[1])-reference_diag["loss"])
        assert errorT <= 1e-10 and errorC <= 1e-12 and error_loss <= 1e-12
        assert diag[2] == steps
        static_comparisons[f"s={scale:g}"] = {"T_max_difference": errorT,
            "C_max_difference": errorC, "loss_difference": error_loss,
            "same_step_count": steps, "independent_dense_reference": True}
    # No boundary exchange and a uniform material field must survive shrinkage.
    uniformT, uniformC = np.full(n, 314.15), np.full(n, .4)
    closedT, closedC, closedlog = advance_interval(uniformT, uniformC, 0., 3600., 30., env,
        faces, centers, volumes, 0., 0., atol, rtol, 60, data, slopes, True, 1., 0.)
    assert np.array_equal(closedT, uniformT) and np.array_equal(closedC, uniformC)
    assert closedlog[1] == 0 and radius_at(3600., data, slopes) < radius_at(0., data, slopes)
    # Save/reload check uses identical time-step boundaries, including total loss.
    actual_env = load_environment(ROOT/"data/attachment1.xlsx")
    args = (actual_env, faces, centers, volumes, h, hm, atol, rtol, 60, data, slopes, True, 1., 0.)
    fullT, fullC, full_diag = advance_interval(T0, C0, 0., 120., .25, *args)
    splitT, splitC, split_diag = advance_interval(T0, C0, 0., 60., .25, *args)
    with tempfile.TemporaryDirectory(prefix="q4_restart_", dir=bootstrap.RUNTIME/"tmp") as tmp:
        path = Path(tmp)/"state.npz"
        np.savez_compressed(path, T=splitT, C=splitC, time=60., loss=split_diag[1])
        with np.load(path, allow_pickle=False) as saved:
            reloadT, reloadC, reload_diag = advance_interval(saved["T"], saved["C"],
                float(saved["time"]), 120., .25, *args)
            reloaded_loss = float(saved["loss"])+reload_diag[1]
    restart = {"T": float(np.max(abs(fullT-reloadT))),
               "C": float(np.max(abs(fullC-reloadC))),
               "loss": float(abs(full_diag[1]-reloaded_loss))}
    assert max(restart.values()) <= LIMITS["restart"]
    balance = float(abs(np.dot(volumes, fullC-C0)/volumes.sum()+full_diag[1]))
    assert balance <= LIMITS["water"]
    assert np.array_equal(T0, 310-2*(centers/RADIUS_0)**2)
    assert np.array_equal(C0, 1.2-.2*(centers/RADIUS_0)**2)
    sample_r = np.array([0., .011, .01198, .012, .02])
    late_time = float(data[-1, 0])
    sampledT, sampledC, mask = sample_fields(fullT, fullC, late_time, actual_env,
        faces, centers, sample_r, h, hm, data, slopes, True, False)
    assert np.array_equal(mask, np.array([True, True, True, False, False]))
    assert np.isnan(sampledT[~mask]).all() and np.isnan(sampledC[~mask]).all()
    surfaceT, surfaceC = surface_fields(fullT, fullC, late_time, actual_env,
        faces, centers, h, hm, data, slopes, True)
    assert abs(sampledT[2]-surfaceT) <= 1e-12 and abs(sampledC[2]-surfaceC) <= 1e-12
    platform_time = 183930.
    platform_radius = radius_at(platform_time, data, slopes)
    assert platform_radius == .012
    platform_samples = np.linspace(0., RADIUS_0, 21)
    platformT, platformC, platform_mask = sample_fields(fullT, fullC, platform_time,
        actual_env, faces, centers, platform_samples, h, hm, data, slopes, True, False)
    platform_surfaceT, platform_surfaceC = surface_fields(fullT, fullC, platform_time,
        actual_env, faces, centers, h, hm, data, slopes, True)
    assert platform_samples[12] == platform_radius and platform_mask[12]
    assert abs(platformT[12]-platform_surfaceT) <= 1e-12
    assert abs(platformC[12]-platform_surfaceC) <= 1e-12
    fakeC = np.full(n, .12); fakeC[n//2] = .2
    maximum = global_maximum(fullT, fakeC, late_time, actual_env,
        faces, centers, h, hm, data, slopes, True)
    assert maximum[0] == .2
    assert abs(maximum[1]-centers[n//2]*.01198/RADIUS_0) <= DOMAIN_ATOL
    return {"passed": True, "radius_source_sha256": hashlib.sha256(radius_path.read_bytes()).hexdigest(),
        "radius_knots_checked": len(data), "exact_plateau_intervals_checked": plateau_intervals,
        "maximum_interpolation_overshoot_m": worst_overshoot,
        "radius_after_last_data_m": float(radius_at(2*data[-1, 0], data, slopes)),
        "fixed_radius_dense_reference": static_comparisons,
        "closed_uniform_shrinking_field_unchanged": True, "closed_loss": float(closedlog[1]),
        "restart": restart, "reference_weight_balance_residual": balance,
        "outside_1_2_cm_blank_at_R_1_198_cm": True, "actual_surface_reconstructed_separately": True,
        "inside_1_2_cm_matches_surface_on_R_1_2_cm_plateau": True,
        "interior_global_maximum_location_scaled": True,
        "interpretation": "Discrete equation, mask and serialization checks; not experimental validation"}


def check_solution(solution, require_completed=True):
    from common.data_loader import load_environment
    from q4.shrinking_radius import radius_at, load_radius_data
    from q4.kernel import advance_interval
    meta = solution["metadata"]
    assert meta["question"] == 4 and meta["model_question"] == 4
    assert meta["threshold"] == .15
    completed = meta.get("completed", meta.get("status") == "complete")
    if require_completed:
        assert completed and meta["status"] == "complete"
    assert float(meta.get("start_time_s", 0.)) == 0
    fields = check_field_domains(solution)
    times, radius = solution["times"], solution["radius_m"]
    last = float(times[-1])
    regular = np.arange(0., math.floor((last+1e-9)/60)*60+1, 60.)
    expected = regular if abs(regular[-1]-last) <= 1e-9 else np.r_[regular, last]
    assert np.array_equal(times, expected)
    data, slopes = solution["radius_data"], solution["radius_slopes"]
    actual_data, actual_slopes = load_radius_data(project_path(meta["radius_path"]))
    assert np.array_equal(data, actual_data) and np.array_equal(slopes, actual_slopes)
    expected_radius = np.asarray([radius_at(float(t), data, slopes, bool(meta["shrinking"])) for t in times])
    assert np.max(abs(radius-expected_radius)) <= DOMAIN_ATOL
    reference_radii = solution["reference_radii"]
    assert np.max(abs(reference_radii-solution["xi"]*RADIUS_0)) <= DOMAIN_ATOL
    assert np.array_equal(solution["reference_physical_radii"], radius[:, None]*solution["xi"][None, :])
    faces, centers, volumes = (solution[k] for k in ("faces", "centers", "volumes"))
    assert len(centers) == meta["cells"] and len(faces) == len(centers)+1
    assert faces[0] == 0 and faces[-1] == RADIUS_0
    assert np.max(abs(volumes-(faces[1:]**2-faces[:-1]**2)/2)) <= 1e-18
    assert np.all(solution["initial_T"] == INITIAL_T) and np.all(solution["initial_C"] == INITIAL_C)
    assert len(solution["initial_C"]) == len(centers)
    assert float(solution["initial_time_s"]) == 0 and float(solution["initial_loss_cumulative"]) == 0
    for key in ("final_T", "final_C"):
        assert solution[key].shape == centers.shape and np.isfinite(solution[key]).all()
        assert np.min(solution[key]) > 0
    scale_final = radius[-1]/RADIUS_0
    assert np.array_equal(solution["final_physical_faces"], scale_final*faces)
    assert np.array_equal(solution["final_physical_centers"], scale_final*centers)
    logs = solution["logs"]
    assert logs.shape == (len(times), 9) and np.isfinite(logs).all()
    assert np.max(logs[:, 7]) <= 1 and np.min(logs[:, 1]) >= -1e-12
    assert np.max(abs(np.cumsum(logs[:, 1])-solution["loss_cumulative"])) <= 1e-10
    assert solution["loss_cumulative"][0] == 0
    expected_balance = solution["mean_C"]+solution["loss_cumulative"]-INITIAL_C
    assert np.max(abs(expected_balance-solution["water_balance"])) <= 1e-12
    water = float(np.max(abs(expected_balance)))
    assert water <= LIMITS["water"]
    mean_final = float(np.dot(volumes, solution["final_C"])/volumes.sum())
    assert abs(mean_final-solution["mean_C"][-1]) <= 1e-12
    assert np.max(abs(logs[1:, 0]-solution["mean_C"][1:])) <= 1e-12
    verified_hashes = {}
    for source_path, expected_hash in meta.get("source_hashes", {}).items():
        path = project_path(source_path)
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        assert actual_hash == expected_hash, f"Changed Q4 input or numerical source: {path}"
        key = path.relative_to(ROOT).as_posix() if path.is_relative_to(ROOT) else path.as_posix()
        verified_hashes[key] = actual_hash
    assert verified_hashes, "Q4 outputs must bind their original data and numerical sources"
    parent_check = None
    if "validation_parent_path" in meta:
        parent_path = project_path(meta["validation_parent_path"])
        parent_hash = hashlib.sha256(parent_path.read_bytes()).hexdigest()
        assert parent_hash == meta["validation_parent_sha256"]
        parent = load_solution(parent_path)
        assert meta["validation_prefix_reused"] is True
        for key in ("times", "temperature_K", "moisture", "valid_mask", "temperature_ref_K",
                    "moisture_ref", "radius_m", "max_C", "max_r", "mean_C", "logs",
                    "loss_cumulative", "water_balance", "surface_moisture", "surface_temperature_K"):
            assert np.array_equal(solution[key][:-1], parent[key][:-1], equal_nan=True), "An event-only study changed its prefix"
        for key in ("initial_T", "initial_C", "faces", "centers", "volumes", "event_anchor_T", "event_anchor_C"):
            assert np.array_equal(solution[key], parent[key])
        assert float(solution["event_anchor_time_s"]) == float(parent["event_anchor_time_s"])
        assert float(solution["event_anchor_loss_cumulative"]) == float(parent["event_anchor_loss_cumulative"])
        parent_check = {"path": meta["validation_parent_path"], "sha256": parent_hash,
                        "prefix_and_native_anchor_identical": True}
    env = load_environment(project_path(meta["environment_path"]))
    assert np.array_equal(solution["environment"], env)
    endpoint_checks = {}
    if completed:
        finish = float(meta["finish_time_s"])
        assert finish == last
        bracket = meta["bracket"]
        assert bracket["g_minus"] >= 0 and bracket["g_plus"] < 0
        assert 0 < bracket["width"] <= meta["epsilon_t"]+1e-10
        assert abs(bracket["t_plus"]-finish) <= 1e-9
        assert abs(bracket["width"]-(bracket["t_plus"]-bracket["t_minus"])) <= 1e-9
        assert np.all(solution["max_C"][:-1] >= .15)
        assert float(solution["event_lower_time_s"]) == bracket["t_minus"]
        assert float(solution["event_upper_time_s"]) == finish
        assert np.array_equal(solution["event_upper_T"], solution["final_T"])
        assert np.array_equal(solution["event_upper_C"], solution["final_C"])
        anchor = float(solution["event_anchor_time_s"])
        anchor_loss = float(solution["event_anchor_loss_cumulative"])
        assert anchor <= bracket["t_minus"] < finish
        for side, Tkey, Ckey in (("minus", "event_lower_T", "event_lower_C"), ("plus", "final_T", "final_C")):
            t = float(bracket[f"t_{side}"])
            R = radius_at(t, data, slopes, bool(meta["shrinking"]))
            extT = float(np.interp(t, env[:, 0], env[:, 1])) if t <= env[-1, 0] else 323.15
            extC = float(np.interp(t, env[:, 0], env[:, 2])) if t <= env[-1, 0] else .05
            independent = independent_boundary_values(solution[Tkey], solution[Ckey], faces, centers,
                R, extT, extC, meta["h"], meta["hm"])
            assert abs(independent["max_C"]-.15-bracket[f"g_{side}"]) <= 1e-12
            assert independent["max_C"] >= .15 if side == "minus" else independent["max_C"] < .15
            T, C, diag = advance_interval(solution["event_anchor_T"], solution["event_anchor_C"],
                anchor, t, meta["event_dt"], env, faces, centers, volumes, meta["h"], meta["hm"],
                meta["atol"], meta["rtol"], meta["max_iterations"], data, slopes,
                bool(meta["shrinking"]), meta["dt_scale"], meta["time_growth"])
            error = max(float(np.max(abs(T-solution[Tkey]))), float(np.max(abs(C-solution[Ckey]))))
            assert error <= LIMITS["restart"]
            loss = float(solution["event_lower_loss_cumulative"]) if side == "minus" else float(solution["loss_cumulative"][-1])
            loss_error = abs(anchor_loss+diag[1]-loss)
            assert loss_error <= LIMITS["restart"]
            endpoint_checks[side] = {**independent, "reintegrated_state_max_difference": error,
                                     "reintegrated_loss_difference": float(loss_error)}
        upper = endpoint_checks["plus"]
        assert abs(upper["max_C"]-solution["max_C"][-1]) <= 1e-12
        assert abs(upper["max_r"]-solution["max_r"][-1]) <= DOMAIN_ATOL
        assert abs(upper["surface_C"]-solution["surface_moisture"][-1]) <= 1e-12
        assert abs(upper["surface_T"]-solution["surface_temperature_K"][-1]) <= 1e-10
        trace = solution["event_trace"]
        assert trace.ndim == 2 and trace.shape[1] == 4 and np.isfinite(trace).all()
        assert np.max(abs(trace[:, 1]-.15-trace[:, 3])) <= 1e-14
        for side in ("minus", "plus"):
            assert np.any((abs(trace[:, 0]-bracket[f"t_{side}"]) < 1e-9)
                          & (abs(trace[:, 3]-bracket[f"g_{side}"]) < 1e-12))
    return {"passed": True, "completed": bool(completed), "field_domains": fields,
        "maximum_water_balance_residual": water, "endpoint_full_state_checks": endpoint_checks,
        "source_hashes_verified": verified_hashes, "initial_state_solved_from_zero": True,
        "fixed_reference_weights": True, "event_parent_check": parent_check,
        "interpretation": "Numerical invariants and original-input consistency; not measured internal-field validation"}


def check_deliverables(solution, excel_path, table_path, endpoint_path=None):
    """Check every fixed physical value, exterior blank, and separate surface."""
    finish = float(solution["metadata"]["finish_time_s"])
    times, radii = solution["times"], solution["radii"]
    requested_times = np.arange(60., math.floor((finish+1e-9)/60)*60+1, 60.)
    rows = matching_indices(times, requested_times)
    cols = matching_indices(radii, np.linspace(0., RADIUS_0, 21), 1e-12)
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    numerical_values, exterior_blanks = 0, 0
    try:
        assert workbook.sheetnames == ["Sheet1"]
        sheet = workbook.worksheets[0]
        assert sheet.max_row == len(requested_times)+1 and sheet.max_column == 23
        iterator = iter(sheet.iter_rows())
        header = next(iterator)
        assert np.allclose([cell.value for cell in header[1:22]], np.arange(21)/10, rtol=0, atol=1e-12)
        assert header[22].value == "药材表面"
        for t, i, row in zip(requested_times, rows, iterator):
            assert row[0].value == t
            for cell, j in zip(row[1:22], cols):
                if solution["valid_mask"][i, j]:
                    assert isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                    assert cell.number_format == "0.0000"
                    assert abs(cell.value-round(float(solution["moisture"][i, j]), 4)) <= 1e-12
                    numerical_values += 1
                else:
                    assert cell.value is None, "Physical exterior must be a blank Excel cell"
                    exterior_blanks += 1
            surface = row[22]
            assert isinstance(surface.value, (int, float)) and not isinstance(surface.value, bool)
            assert surface.number_format == "0.0000"
            assert abs(surface.value-round(float(solution["surface_moisture"][i]), 4)) <= 1e-12
        assert next(iterator, None) is None
    finally:
        workbook.close()
    with Path(table_path).open(encoding="utf-8-sig", newline="") as stream:
        table = list(csv.DictReader(stream))
    table_times = np.arange(21600., math.floor((finish+1e-9)/21600)*21600+1, 21600.)
    if not len(table_times) or abs(table_times[-1]-finish) > 1e-9:
        table_times = np.r_[table_times, finish]
    assert len(table) == len(table_times)
    physical_positions = (0., .5, 1., 1.5, 2.)
    table_cols = matching_indices(radii, np.asarray(physical_positions)/100, 1e-12)
    for index, (row, t) in enumerate(zip(table, table_times)):
        i = int(matching_indices(times, np.array([t]))[0])
        assert abs(float(row["time_s"])-t) <= 1e-9
        assert abs(float(row["time_h"])-t/3600) <= 1e-12
        assert abs(float(row["radius_cm"])-100*solution["radius_m"][i]) <= 1e-12
        assert row["row_kind"] == ("endpoint" if index == len(table)-1 else "regular_6h")
        for r, j in zip(physical_positions, table_cols):
            expected = f"{solution['moisture'][i, j]:.4f}" if solution["valid_mask"][i, j] else ""
            assert row[f"r={r:g} cm"] == expected
        assert row["surface_C"] == f"{solution['surface_moisture'][i]:.4f}"
    if endpoint_path is not None:
        endpoint = json.loads(Path(endpoint_path).read_text(encoding="utf-8"))
        assert abs(endpoint["finish_time_s"]-finish) <= 1e-9
        assert abs(endpoint["finish_time_h"]-finish/3600) <= 1e-12
        assert abs(endpoint["radius_m"]-solution["radius_m"][-1]) <= 1e-14
        assert endpoint["bracket"] == solution["metadata"]["bracket"]
        for key in ("max_C", "max_r", "surface_moisture", "mean_C", "loss_cumulative", "water_balance"):
            assert abs(endpoint[key]-float(solution[key][-1])) <= 1e-14
        assert endpoint["max_C"] < .15 and abs(endpoint["g_plus"]-(endpoint["max_C"]-.15)) <= 1e-14
        assert np.array_equal(np.asarray(endpoint["radii_m"]), radii)
        assert np.array_equal(np.asarray(endpoint["valid_mask"]), solution["valid_mask"][-1])
        assert len(endpoint["moisture"]) == len(radii)
        for value, original, valid in zip(endpoint["moisture"], solution["moisture"][-1], solution["valid_mask"][-1]):
            assert value == original if valid else value is None
    return {"passed": True, "excel_regular_rows_checked": len(requested_times),
        "excel_physical_C_values_checked": numerical_values,
        "excel_exterior_blank_cells_checked": exterior_blanks,
        "excel_separate_surface_values_checked": len(requested_times),
        "table6_rows_checked": len(table), "exact_endpoint_checked": endpoint_path is not None,
        "interpretation": "Artifact-to-array identity with the moving-domain mask, not physical accuracy"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--core", action="store_true")
    parser.add_argument("--solution", type=Path, default=ROOT/"q4/solution.npz")
    parser.add_argument("--skip-deliverables", action="store_true")
    parser.add_argument("--output", type=Path, default=ROOT/"validation/q4_checks.json")
    args = parser.parse_args()
    if args.core:
        report = {"core": core_checks()}
    else:
        solution = load_solution(args.solution)
        report = {"solution": check_solution(solution)}
        if not args.skip_deliverables:
            report["deliverables"] = check_deliverables(solution, ROOT/"submission_results/result4.xlsx",
                ROOT/"q4/tables/table6.csv", ROOT/"q4/endpoint.json")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
