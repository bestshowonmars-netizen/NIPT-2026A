"""Compare saved Q3 runs and spreadsheets; never launch a numerical solve.

This verifies a self-contained reproduction of the same algorithm, not an
independent physical model or validation against measured internal fields.
"""
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import bootstrap
import argparse
import hashlib
import json
import math
import numpy as np
from openpyxl import load_workbook
from validation.evidence_paths import project_path

ROOT = bootstrap.ROOT
INTERPRETATION = (
    "Self-contained reproduction of the same discretization and algorithm, "
    "including a fresh Q2 solve from the original attachment and initial state; "
    "not independent physical validation or comparison with measured fields."
)


def local_path(value):
    """Resolve recorded project paths inside the current checkout after copying."""
    return project_path(value, ROOT)


def receipt_digest(recorded_hashes, path):
    matches = [(original, expected) for original, expected in recorded_hashes.items()
               if local_path(original).resolve() == Path(path).resolve()]
    assert matches, f"Fresh-solve receipt does not identify {path}"
    actual = digest(path)
    assert all(expected == actual for _, expected in matches)
    return [original for original, _ in matches]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_npz(path):
    with np.load(path, allow_pickle=False) as saved:
        result = {key: saved[key].copy() for key in saved.files}
    if "metadata_json" in result:
        result["metadata"] = json.loads(str(result.pop("metadata_json")))
    return result


def difference(left, right, tolerance):
    left, right = np.asarray(left), np.asarray(right)
    assert left.shape == right.shape, (left.shape, right.shape)
    assert np.isfinite(left).all() and np.isfinite(right).all()
    error = float(np.max(np.abs(left-right))) if left.size else 0.0
    assert error <= tolerance, (error, tolerance)
    return {"shape": list(left.shape), "array_equal": bool(np.array_equal(left, right)),
            "max_abs_difference": error, "tolerance": tolerance}


def coordinates(values, requested, tolerance):
    """Select existing coordinates only; interpolation is forbidden here."""
    values, requested = np.asarray(values), np.asarray(requested)
    assert values.ndim == 1 and len(values) and np.all(np.diff(values) > 0)
    indices = np.minimum(np.searchsorted(values, requested), len(values)-1)
    previous = np.maximum(indices-1, 0)
    take_previous = abs(values[previous]-requested) < abs(values[indices]-requested)
    indices[take_previous] = previous[take_previous]
    assert np.all(abs(values[indices]-requested) <= tolerance), "Missing comparison coordinates"
    return indices


def compare_prefix(fresh, main_q2, main_final):
    assert fresh["metadata"]["question"] == 2 and fresh["metadata"]["cells"] == 960
    assert float(fresh["times"][-1]) == float(main_final["time_s"]) == 10800.0
    mapping = {"final_T": "temperature_K", "final_C": "moisture",
               "faces": "faces", "centers": "centers", "volumes": "volumes"}
    native = {key: difference(fresh[key], main_final[target], 1e-12)
              for key, target in mapping.items()}
    required = np.linspace(0.0, .02, 21)
    a = coordinates(fresh["radii"], required, 1e-12)
    b = coordinates(main_q2["radii"], required, 1e-12)
    history = {key: difference(fresh[key][:, a], main_q2[key][:, b],
                               1e-10 if key == "temperature_K" else 1e-12)
               for key in ("temperature_K", "moisture")}
    return {"native_grid_arrays": native, "every_second_21_point_history": history,
            "times": difference(fresh["times"], main_q2["times"], 0),
            "logs": difference(fresh["logs"], main_q2["logs"], 0),
            "fresh_solve_metadata": fresh["metadata"]}


def spreadsheet_values(path, expected_times):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        sheet = workbook.worksheets[0]
        assert sheet.max_row == len(expected_times)+1 and sheet.max_column == 22
        rows = iter(sheet.iter_rows())
        header = next(rows)
        difference([cell.value for cell in header[1:]], np.arange(21)/10, 1e-12)
        result = np.empty((len(expected_times), 21))
        for index, (expected_time, row) in enumerate(zip(expected_times, rows)):
            assert row[0].value == expected_time
            for j, cell in enumerate(row[1:]):
                assert isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                assert cell.number_format == "0.0000"
                result[index, j] = cell.value
        assert next(rows, None) is None
        return result
    finally:
        workbook.close()


def compare_existing(main_path, independent_path, main_excel, independent_excel,
                     prefix_path, q2_path, final_path, prefix_evidence, attachment):
    main_path, independent_path, main_excel, independent_excel, q2_path, final_path, attachment = (
        local_path(path) for path in (main_path, independent_path, main_excel,
                                     independent_excel, q2_path, final_path, attachment))
    main, independent = load_npz(main_path), load_npz(independent_path)
    mm, im = main["metadata"], independent["metadata"]
    prefix_path = local_path(prefix_path if prefix_path is not None else im["prefix_path"])
    fresh, q2, final = load_npz(prefix_path), load_npz(q2_path), load_npz(final_path)
    assert mm["status"] == "complete" and mm["completed"] is True
    assert mm["cells"] == 960 and mm["model_question"] == 2
    assert im["standalone_same_algorithm"] is True
    assert local_path(im["prefix_path"]).resolve() == prefix_path.resolve()
    assert im["prefix_sha256"] == digest(prefix_path)
    assert im["source_attachment_sha256"] == digest(attachment)
    assert isinstance(im["from_initial"], bool)
    evidence = None
    if prefix_evidence is None and not im["from_initial"]:
        candidates = (prefix_path.parent/"prefix_comparison.json",
                      Path(independent_path).parent/"prefix_comparison.json")
        prefix_evidence = next((path for path in candidates if path.is_file()), None)
    if prefix_evidence is not None:
        prefix_evidence = local_path(prefix_evidence)
        evidence = json.loads(prefix_evidence.read_text(encoding="utf-8"))
        assert evidence["from_original_attachment_and_initial_state"] is True
        assert evidence["main_cache_used_by_independent_solver"] is False
        assert evidence["fresh_solve_process_exit_code"] == 0
        evidence_hashes = evidence["files_sha256"]
        receipt_digest(evidence_hashes, prefix_path)
        receipt_digest(evidence_hashes, attachment)
    elif not im["from_initial"]:
        raise AssertionError("A reused standalone prefix requires its fresh-solve receipt")
    assert fresh["metadata"]["initial_T"] == 301.15 and fresh["metadata"]["initial_C"] == 2.55
    assert np.all(fresh["temperature_K"][0] == 301.15) and np.all(fresh["moisture"][0] == 2.55)
    for key in ("dt_max", "event_dt"):
        assert mm[key] == im[key]
    assert mm["epsilon_t"] == im["event_tol"]
    finish_main, finish_independent = float(mm["finish_time_s"]), float(im["finish_time_s"])
    finish_difference = abs(finish_main-finish_independent)
    assert finish_difference <= 1e-9
    regular = np.arange(0., math.floor((finish_main+1e-9)/60)*60+1, 60.)
    expected = regular if regular[-1] == finish_main else np.r_[regular, finish_main]
    assert np.array_equal(main["times"], expected)
    assert np.array_equal(independent["times"], expected)
    required = np.linspace(0., .02, 21)
    main_cols = coordinates(main["radii"], required, 1e-12)
    independent_cols = coordinates(independent["radii"], required, 1e-12)
    main_rows = coordinates(main["times"], regular, 1e-9)
    independent_rows = coordinates(independent["times"], regular, 1e-9)
    regular_fields, all_fields = {}, {}
    for key in ("temperature_K", "moisture"):
        tolerance = 1e-10 if key == "temperature_K" else 1e-12
        regular_fields[key] = difference(main[key][main_rows][:, main_cols],
            independent[key][independent_rows][:, independent_cols], tolerance)
        all_fields[key] = difference(main[key][:, main_cols],
            independent[key][:, independent_cols], tolerance)
    native = {key: difference(main[key], independent[key], 1e-12)
              for key in ("final_T", "final_C", "faces", "centers", "volumes")}
    post = main["times"] >= 10800
    maxima = {key: difference(main[key][post], independent[key][post], 1e-12)
              for key in ("max_C", "max_r")}
    bracket_errors = {}
    for label, metadata in (("main", mm), ("independent", im)):
        bracket = metadata["bracket"]
        assert bracket["g_minus"] >= 0 and bracket["g_plus"] < 0
        assert 0 < bracket["width"] <= im["event_tol"]+1e-10
        assert abs(bracket["t_plus"]-metadata["finish_time_s"]) <= 1e-9
        assert abs(bracket["width"]-(bracket["t_plus"]-bracket["t_minus"])) <= 1e-9
    for key in ("t_minus", "t_plus", "g_minus", "g_plus", "width"):
        error = abs(mm["bracket"][key]-im["bracket"][key])
        assert error <= (1e-12 if key.startswith("g_") else 1e-9)
        bracket_errors[key] = error
    trace = difference(main["event_trace"], independent["event_trace"], 1e-12)
    assert float(independent["max_C"][-1]) < .15
    balance = difference(main["water_balance"], independent["water_balance"], 1e-10)
    water_max = float(np.max(abs(independent["water_balance"])))
    assert water_max <= 1e-8
    assert abs(water_max-im["max_water_balance"]) <= 1e-14
    excel_times = regular[regular > 0]
    main_values = spreadsheet_values(main_excel, excel_times)
    independent_values = spreadsheet_values(independent_excel, excel_times)
    workbook_comparison = difference(main_values, independent_values, 1e-12)
    workbook_expectations = {}
    for label, solution, columns, values in (
        ("main", main, main_cols, main_values),
        ("independent", independent, independent_cols, independent_values),
    ):
        rows = coordinates(solution["times"], excel_times, 1e-9)
        expected_values = np.asarray([[round(float(value), 4) for value in row]
                                     for row in solution["moisture"][rows][:, columns]])
        workbook_expectations[label] = difference(values, expected_values, 1e-12)
    files = [main_path, independent_path, main_excel, independent_excel, prefix_path,
             q2_path, final_path, attachment,
             ROOT/"paper_appendix_minimal/minimal_solver.py",
             ROOT/"paper_appendix_minimal/minimal_q3.py", Path(__file__)]
    if prefix_evidence is not None:
        files.append(prefix_evidence)
    return {"passed": True, "interpretation": INTERPRETATION,
        "fresh_q2_source": {"prefix_path": str(Path(prefix_path).resolve()),
            "recorded_prefix_path": im["prefix_path"],
            "prefix_sha256": digest(prefix_path),
            "provenance_mode": "single_standalone_from_initial_run" if im["from_initial"] else "fresh_prefix_receipt",
            "standalone_from_initial": im["from_initial"],
            "evidence_path": str(prefix_evidence.resolve()) if prefix_evidence is not None else None,
            "fresh_solve_command": evidence["fresh_solve_command"] if evidence is not None else None,
            "standalone_continuation_reused_independently_computed_prefix": not im["from_initial"],
            "fresh_q2_comparison": compare_prefix(fresh, q2, final)},
        "configuration": {"cells": 960, "dt_max": mm["dt_max"],
            "event_dt": mm["event_dt"], "epsilon_t": mm["epsilon_t"]},
        "all_regular_60_second_21_point_fields": regular_fields,
        "regular_values_checked_per_field": int(len(regular)*21),
        "regular_time_range_s": [float(regular[0]), float(regular[-1])],
        "all_outputs_including_exact_endpoint": all_fields,
        "full_960_cell_terminal_state_and_geometry": native,
        "full_grid_maximum_history_after_3h": maxima,
        "endpoint": {"main_time_s": finish_main, "independent_time_s": finish_independent,
            "time_difference_s": finish_difference, "main_bracket": mm["bracket"],
            "independent_bracket": im["bracket"], "bracket_abs_differences": bracket_errors,
            "event_trace": trace, "strict_independent_final_max_C": float(independent["max_C"][-1])},
        "excel": {"data_rows_checked": len(excel_times),
            "all_C_values_checked_per_workbook": int(main_values.size),
            "between_workbooks": workbook_comparison,
            "each_workbook_against_own_unrounded_npz": workbook_expectations},
        "water_balance": {"independent_max_abs_residual": water_max,
            "between_runs": balance,
            "meaning": "Residual accumulated within the standalone full Q2+Q3 numerical trajectory"},
        "files_sha256": {str(Path(path).resolve()): digest(path) for path in files}}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--main", type=Path, default=ROOT/"q3/solution.npz")
    parser.add_argument("--independent", type=Path, default=ROOT/
                        "paper_appendix_minimal/reproduced/q3/q3_independent.npz")
    parser.add_argument("--main-excel", type=Path, default=ROOT/"submission_results/result3.xlsx")
    parser.add_argument("--independent-excel", type=Path, default=ROOT/
                        "paper_appendix_minimal/reproduced/q3/result3.xlsx")
    parser.add_argument("--prefix", type=Path,
                        help="Default: actual prefix_path recorded in the independent NPZ")
    parser.add_argument("--prefix-evidence", type=Path,
                        help="Default: prefix_comparison.json beside a reused prefix; unnecessary for --from-initial runs")
    parser.add_argument("--q2", type=Path, default=ROOT/"q2/solution.npz")
    parser.add_argument("--q2-final", type=Path, default=ROOT/"q2/final_state.npz")
    parser.add_argument("--attachment", type=Path, default=ROOT/"data/attachment1.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT/"validation/q3_reproduction.json")
    args = parser.parse_args(argv)
    report = compare_existing(args.main, args.independent, args.main_excel, args.independent_excel,
        args.prefix, args.q2, args.q2_final, args.prefix_evidence, args.attachment)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
    return report


if __name__ == "__main__":
    main()
