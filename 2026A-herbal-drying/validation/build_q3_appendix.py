"""Build the appendix Q3 source from reviewed numerical functions.

The emitted program imports only the local standalone Q1/Q2 appendix solver,
never common/q1/q2/q3/run_all. This is a same-algorithm reproduction artifact.
"""
from pathlib import Path
import ast

ROOT = Path(__file__).resolve().parents[1]
core = ast.parse((ROOT / "common/continuation.py").read_text(encoding="utf-8"))
event = ast.parse((ROOT / "q3/end_time_detector.py").read_text(encoding="utf-8"))
functions = "\n\n".join(ast.unparse(node) for tree in (core, event) for node in tree.body
                        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)))
header = '''"""Standalone question-three continuation; see appendix README for inputs."""
from pathlib import Path
import argparse, json, time, hashlib
import minimal_solver as base
import numpy as np
from numba import njit
from openpyxl import load_workbook
from minimal_solver import (fill_properties, links, linear_step, residual,
                            reconstruct, environment_at, load_environment)
'''
driver = r'''

def main(argv=None):
    parser=argparse.ArgumentParser(description="Standalone Q3 same-algorithm reproduction")
    folder=Path(__file__).resolve().parent
    parser.add_argument("--prefix",type=Path,default=folder.parent/"q2/solution.npz")
    parser.add_argument("--from-initial",action="store_true",
                        help="Recompute the full Q2 prefix using the standalone appendix kernel")
    parser.add_argument("--attachment",type=Path,default=folder.parent/"data/attachment1.xlsx")
    parser.add_argument("--template",type=Path,default=folder.parent/"data/templates/result3_template.xlsx")
    parser.add_argument("--output",type=Path,default=folder/"reproduced/q3")
    parser.add_argument("--dt-max",type=float,default=0.03125)
    parser.add_argument("--event-dt",type=float,default=0.05)
    parser.add_argument("--event-tol",type=float,default=0.01)
    args=parser.parse_args(argv)
    if min(args.dt_max,args.event_dt,args.event_tol)<=0:
        raise ValueError("Numerical scales must be positive")
    args.output.mkdir(parents=True,exist_ok=True)
    environment=load_environment(args.attachment)
    config=dict(question=2,end_time=10800,radius=.02,h=25.,hm=8e-7,
                atol=1e-10,rtol=1e-11,max_iterations=60,initial_T=301.15,
                initial_C=2.55,time_growth=600.)
    if args.from_initial:
        prefix=base.solve(config,environment,960,1/768,sample_count=21,progress=True)
        prefix_path=args.output/"q2_prefix_independent.npz"
        base.save_solution(prefix,prefix_path)
    else:
        prefix_path=args.prefix
        prefix=base.load_solution(args.prefix)
        assert prefix["metadata"]["question"]==2 and prefix["times"][-1]==10800.
        assert prefix["metadata"]["cells"]==960
        for key in ["h","hm","initial_T","initial_C"]:
            assert prefix["metadata"][key]==config[key]
    config.update(threshold=.15,event_dt=min(args.event_dt,args.dt_max),
                  epsilon_t=args.event_tol)
    faces,centers,V=prefix["faces"],prefix["centers"],prefix["volumes"]
    radii=np.linspace(0,.02,21)
    stride=(len(prefix["radii"])-1)//20
    assert np.allclose(prefix["radii"][::stride],radii,rtol=0,atol=1e-15)
    times=list(prefix["times"][::60])
    fieldsT=list(prefix["temperature_K"][::60,::stride])
    fieldsC=list(prefix["moisture"][::60,::stride])
    maxima=[float("nan")]*len(times)
    argmax=[float("nan")]*len(times)
    T,C=prefix["final_T"].copy(),prefix["final_C"].copy()
    t=10800.
    m,r,_,_=global_maximum(T,C,t,environment,faces,centers,25.,8e-7)
    maxima[-1],argmax[-1]=m,r
    loss=float(prefix["logs"][:,1].sum())
    prefix_cumulative=np.cumsum(prefix["logs"][:,1])
    balances=[0.]+list(prefix["logs"][59::60,0]+prefix_cumulative[59::60]-2.55)
    trace=[]
    last_print=time.perf_counter()
    while t<30*24*3600:
        next_time=t+60.
        tn,cn,diag=advance_interval(T,C,t,next_time,args.dt_max,environment,
                                   faces,centers,V,25.,8e-7,1e-10,1e-11,60)
        m,r,_,_=global_maximum(tn,cn,next_time,environment,faces,centers,25.,8e-7)
        terminal=False
        if m<.15:
            event=refine_endpoint(T,C,t,next_time,environment,faces,centers,V,config)
            trace.extend(event["trace"].tolist())
            tn,cn,diag=event["T"],event["C"],event["diag"]
            next_time,m,r=event["time"],event["M"],event["r"]
            terminal=event["crossed"]
        T,C,t=tn,cn,next_time
        loss+=float(diag[1])
        st,sc=sample_fields(T,C,t,environment,faces,centers,radii,25.,8e-7)
        times.append(t);fieldsT.append(st);fieldsC.append(sc)
        maxima.append(m);argmax.append(r)
        balances.append(float(np.dot(V,C)/V.sum()+loss-2.55))
        if terminal:
            bracket=event["bracket"]
            break
        if time.perf_counter()-last_print>30:
            print(f"Q3 appendix: t={t/3600:.3f} h, max C={m:.6f}",flush=True)
            last_print=time.perf_counter()
    else:
        raise RuntimeError("No endpoint by the computational safety bound; no success reported")
    assert bracket["g_minus"]>=0 and bracket["g_plus"]<0
    meta={"finish_time_s":t,"finish_time_h":t/3600,"bracket":bracket,
          "dt_max":args.dt_max,"event_dt":config["event_dt"],"event_tol":args.event_tol,
          "from_initial":args.from_initial,"standalone_same_algorithm":True,
          "prefix_path":str(prefix_path.resolve()),
          "prefix_sha256":hashlib.sha256(prefix_path.read_bytes()).hexdigest(),
          "source_attachment_sha256":hashlib.sha256(args.attachment.read_bytes()).hexdigest(),
          "max_water_balance":float(np.max(np.abs(balances)))}
    np.savez_compressed(args.output/"q3_independent.npz",
        times=np.asarray(times),radii=radii,temperature_K=np.asarray(fieldsT),
        moisture=np.asarray(fieldsC),max_C=np.asarray(maxima),max_r=np.asarray(argmax),
        final_T=T,final_C=C,faces=faces,centers=centers,volumes=V,
        water_balance=np.asarray(balances),event_trace=np.asarray(trace),
        metadata_json=json.dumps(meta,ensure_ascii=False))
    wb=load_workbook(args.template)
    ws=wb.worksheets[0]
    for row in ws:
        for cell in row:
            if cell.coordinate!="A1":cell.value=None
    for j in range(21):ws.cell(1,j+2,round(j*.1,1))
    output_row=2
    for ti,ci in zip(times,fieldsC):
        if ti<=0 or abs(ti/60-round(ti/60))>1e-9:continue
        ws.cell(output_row,1,int(round(ti)))
        for j,value in enumerate(ci):
            cell=ws.cell(output_row,j+2,round(float(value),4))
            cell.number_format="0.0000"
        output_row+=1
    ws.freeze_panes="B2"
    wb.properties.creator="";wb.properties.lastModifiedBy=""
    wb.save(args.output/"result3.xlsx")
    (args.output/"endpoint.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(meta,ensure_ascii=False,indent=2),flush=True)
    return meta

if __name__=="__main__":
    main()
'''
target=ROOT/"paper_appendix_minimal/minimal_q3.py"
target.write_text(header+"\n\n"+functions+"\n"+driver,encoding="utf-8")
print(target)
