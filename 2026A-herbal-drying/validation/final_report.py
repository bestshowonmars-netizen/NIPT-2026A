"""Assemble checks from saved computations and verify independent appendix outputs."""
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bootstrap,json,hashlib,ast
import numpy as np
from openpyxl import load_workbook
from common.solver import load_solution
from common.plotting import figure3,convergence_figure
ROOT=bootstrap.ROOT

def main():
    records=json.loads((ROOT/'validation/strict_convergence.json').read_text(encoding='utf-8'))
    assert all(item['passed'] for item in records if item.get('required',True)), 'Strict refinement targets are not all satisfied'
    sys.path.insert(0,str(ROOT/'paper_appendix_minimal'))
    import minimal_solver as independent
    tree=ast.parse((ROOT/'paper_appendix_minimal/minimal_solver.py').read_text(encoding='utf-8'))
    for node in ast.walk(tree):
        if isinstance(node,ast.ImportFrom):assert not (node.module or '').startswith(('common','q1','q2','run_all'))
        if isinstance(node,ast.Import):assert all(not name.name.startswith(('common','q1','q2','run_all')) for name in node.names)
    replication=[]
    for q in [1,2]:
        full=load_solution(ROOT/f'q{q}/solution.npz')
        small=independent.load_solution(ROOT/f'paper_appendix_minimal/reproduced/q{q}_independent.npz')
        result={'question':q}
        for key in ['temperature_K','moisture']:
            error=float(np.max(abs(full[key][:,::5]-small[key])))
            assert error<1e-10
            result[key+'_max_difference']=error
        for key in ['final_T','final_C','faces','centers','volumes']:
            assert np.array_equal(full[key],small[key])
        out=ROOT/f'paper_appendix_minimal/reproduced/result{q}.xlsx'
        independent.export_result(small,ROOT/f'data/templates/result{q}_template.xlsx',out)
        mainbook=load_workbook(ROOT/f'submission_results/result{q}.xlsx',read_only=True,data_only=True)
        smallbook=load_workbook(out,read_only=True,data_only=True)
        assert mainbook.sheetnames==smallbook.sheetnames
        checked=0
        for title in mainbook.sheetnames:
            a,b=mainbook[title],smallbook[title]
            assert a.max_row==b.max_row and a.max_column==b.max_column
            for row_a,row_b in zip(a.values,b.values):
                assert row_a==row_b
                checked+=len(row_a)
        mainbook.close();smallbook.close()
        result['all_excel_cells_identical']=True
        result['cells_compared_including_headers']=checked
        result['independent_of_main_modules']=True
        replication.append(result)
    (ROOT/'validation/appendix_reproduction.json').write_text(json.dumps(replication,indent=2),encoding='utf-8')
    originals=ROOT.parent.parent/'CUMCM2026Problems/A题/附件'
    manifest=[]
    for source,target in [('附件1.xlsx','data/attachment1.xlsx'),('附件3/result1.xlsx','data/templates/result1_template.xlsx'),('附件3/result2.xlsx','data/templates/result2_template.xlsx')]:
        a=hashlib.sha256((originals/source).read_bytes()).hexdigest()
        b=hashlib.sha256((ROOT/target).read_bytes()).hexdigest()
        assert a==b
        manifest.append({'source':source,'project_copy':target,'sha256':a})
    (ROOT/'data/input_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    checks=json.loads((ROOT/'validation/result_checks.json').read_text(encoding='utf-8'))
    solutions=[load_solution(ROOT/f'q{q}/solution.npz') for q in [1,2]]
    # Every required output radius of the dense plotting run matches the audited production run.
    from validation.grid_convergence import run
    dense_checks=[]
    for q,sol in zip([1,2],solutions):
        audited=run(q,960,1/768,growth=600)
        d={key:float(np.max(abs(sol[key][:,::5]-audited[key]))) for key in ['temperature_K','moisture']}
        assert max(d.values())<1e-10
        dense_checks.append({'question':q,**d})
    summary={'strict_refinement_passed':True,'required_comparisons':sum(r.get('required',True) for r in records),
             'full_result_values_checked':sum(r['all_excel_values_checked'] for r in checks['results']),
             'appendix_reproduction':replication,'dense_sampling_agreement':dense_checks,
             'input_copies_match_sources':True,'question_count':2}
    (ROOT/'validation/final_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# 精度与复现报告','',
      '本报告对应正式交付的第一、二问。所有必要的空间、时间、共同加密和Picard容差比较均通过预定门槛；未通过的粗档记录单独保留。',
      '', '## 正式配置','',
      '| 项目 | 第一问 | 第二问 |','|---|---:|---:|',
      '| 范围 / s | 0—1800 | 0—10800 |','| 径向单元数 | 960 | 960 |',
      '| 时间尺度 τ / s | 1/768 | 1/768 |','| 时间步上限 / s | 1/192 | 1/192 |',
      '| Picard绝对/相对容差 | 1e-10 / 1e-11 | 1e-10 / 1e-11 |']
    for title,key,fmt in [('接受步数','steps','d'),('Picard总迭代数','iterations','d'),('回退步数','rejections','d'),
        ('阻尼次数','damped_iterations','d'),('最大缩放残差','max_scaled_residual','.3e'),('最大归一化水量残差','max_water_balance_residual','.3e')]:
        values=[format(s['metadata'][key],fmt) for s in solutions]
        lines.append('| '+title+' | '+' | '.join(values)+' |')
    lines += ['', '计算步长在下一整数秒处截断，初期还受τ(1+t/600)²约束；输出始终为每1秒。完整网格与末状态已经保存。',
      '', '## 严格加密结果','',
      '比较覆盖每个题目输出整数秒与21个径向位置，包含中心和表面。表中温度差单位为°C（与K差值相同），含水率差单位为kg/kg。空间比较固定τ=1/1024秒；时间比较固定1280单元。',
      '', '| 问题 | 检查 | 比较 | 最大温度差 | 最大含水率差 | 门槛 | 判定 |', '|---|---|---|---:|---:|---:|---|']
    names={'space':'空间','time':'时间','joint':'共同加密','picard':'Picard容差'}
    for record in records:
        a,b=record['coarse'],record['fine'];kind=record['kind']
        if kind=='space':desc=f"N={a['N']}→{b['N']}"
        elif kind=='time':desc=f"τ={a['dt']:.9g}→{b['dt']:.9g} s"
        elif kind=='joint':desc=f"N={a['N']}→{b['N']}；τ=1/768→1/1024"
        else:desc='绝对及相对容差同时缩小100倍'
        verdict='达到' if record['passed'] else '未达到；保留的粗档'
        lines.append(f"| {record['q']} | {names[kind]} | {desc} | {record['T']['maximum']:.6e} | {record['C']['maximum']:.6e} | {record['limit']:.0e} | {verdict} |")
    lines += ['', '最终空间验收使用连续两次960→1280→1600的加密；时间验收使用连续两次τ=1/512→1/768→1/1024的加密。均要求两个场的最大差异小于1e-5；共同加密小于2e-5；加严Picard容差后的差异小于1e-6。原640→960空间比较未达到单项门槛，因此继续加密至1600，没有将失败记录计为通过。',
      '', '上述检查满足本次约定的四位小数数值稳定性验收。更细计算仍是数值参考，不是解析真解；不能由此推出未知精确解的严格误差上界，也不保证处在四舍五入分界附近的每个末位字符永远不变。实际舍入仅用于最终展示，原始高精度数组保留。',
      '', '最大偏差的具体秒数和径向位置均记录在strict_convergence.json中，缩放方程残差、每秒步数与质量核算见各问tables/solver_diagnostics.csv。',
      '', '## 物理与文件检查','',
      '- 已检查全部241个环境观测节点、插值中点、K/°C转换和预定延拓约定。',
      '- 两问物性公式、中心对称二次重构、Robin表面通量、均匀平衡解和独立稠密线性系统对照均通过。',
      f"- 封闭边界非均匀扩散的归一化水量偏差为{checks['core']['closed_boundary_mass_error']:.3e}。正式两问的归一化水量残差均小于1e-8。",
      '- 已逐格检查两份Excel共529200个结果数值及对应表头、时空采样、数值类型和四位小数格式。表1—4与完整结果逐项一致。',
      '- 两问各自的0秒高精度状态均严格采用题设均匀初态；第二问没有拼接第一问。',
      '- 附件1和两个模板副本与原文件SHA-256一致，记录在data/input_manifest.json。',
      '', '## 独立附录复现','',
      '附录代码在自己的目录中单独启动Python进程，未导入主项目求解模块，重新计算两问完整范围。输入、网格、时间步规则和容差与正式计算一致。',
      '', '| 问题 | 高精度温度最大差 | 高精度含水率最大差 | 末网格状态 | Excel全部单元格 |', '|---|---:|---:|---|---|']
    for rep in replication:
        lines.append(f"| {rep['question']} | {rep['temperature_K_max_difference']:.3e} | {rep['moisture_max_difference']:.3e} | 逐元素完全相同 | 完全相同 |")
    lines += ['', '复现文件位于paper_appendix_minimal/reproduced。图用101个径向采样位置，题目21个位置是其子集；该子集与验收计算及独立附录计算的一致性也已检查。',
      '', '## 结果范围说明','',
      '本项目仅给出第一问前1800秒、第二问前3小时的预测。没有第三、四问干燥终点、收缩模型或对应Excel；没有把数值一致性、水量守恒或环境输入插值一致性称为内部温湿场实验验证。']
    (ROOT/'validation/精度与复现报告.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    figure3(solutions[0],ROOT/'q1/figures')
    convergence_figure(records,ROOT/'validation/figures')
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
    print('Final endpoints:')
    for q,s in zip([1,2],solutions):
        print(q,s['temperature_K'][-1,[0,-1]]-273.15,s['moisture'][-1,[0,-1]])
if __name__=='__main__':main()
