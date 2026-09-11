"""Standalone question-three continuation; see appendix README for inputs."""
from pathlib import Path
import argparse, json, time, hashlib
import minimal_solver as base
import numpy as np
from numba import njit
from openpyxl import load_workbook
from minimal_solver import (fill_properties, links, linear_step, residual,
                            reconstruct, environment_at, load_environment)


@njit(cache=True)
def advance_interval(T0, C0, t_start, t_end, dt_max, env, faces, centers, volume, h, hm, atol, rtol, max_iterations):
    """Return new full-grid states and 9 diagnostics; never mutate input state.

    Diagnostic order agrees with integrate_block:
    mean_C, interval_loss, steps, iterations, rejections, min_dt, max_dt,
    worst_scaled_residual, damped_iterations.
    """
    if t_end < t_start or dt_max <= 0.0 or max_iterations < 1:
        raise ValueError('Invalid continuation interval or step configuration')
    T, C = (T0.copy(), C0.copy())
    n = len(T)
    a, k, D = (np.empty(n), np.empty(n), np.empty(n))
    ones = np.ones(n)
    gt, gc = (np.empty(n + 1), np.empty(n + 1))
    work = np.empty((2, n))
    guessT, guessC = (np.empty(n), np.empty(n))
    nextT, nextC = (np.empty(n), np.empty(n))
    total_volume = np.sum(volume)
    t, dt_limit = (t_start, dt_max)
    loss, steps, iterations, rejected = (0.0, 0, 0, 0)
    minimum, maximum, worst, damped = (1e+100, 0.0, 0.0, 0)
    while t < t_end - 1e-12:
        knot_index = np.searchsorted(env[:, 0], t + 1e-10, side='right')
        next_knot = env[knot_index, 0] if knot_index < len(env) else t_end
        dt = min(dt_limit, t_end - t, next_knot - t)
        extT, extC = environment_at(t + dt, env)
        guessT[:], guessC[:] = (T, C)
        converged, previous_change, omega = (False, 1e+100, 1.0)
        for it in range(max_iterations):
            fill_properties(guessT, guessC, 2, a, k, D)
            links(k, h, faces, centers, gt)
            linear_step(T, a, volume, gt, extT, dt, nextT, work)
            fill_properties(nextT, guessC, 2, a, k, D)
            links(D, hm, faces, centers, gc)
            linear_step(C, ones, volume, gc, extC, dt, nextC, work)
            if not np.all(np.isfinite(nextT)) or not np.all(np.isfinite(nextC)):
                break
            if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                omega *= 0.5
            change = 0.0
            for j in range(n):
                change = max(change, abs(nextT[j] - guessT[j]) / (atol + rtol * abs(nextT[j])))
                change = max(change, abs(nextC[j] - guessC[j]) / (atol + rtol * abs(nextC[j])))
            if change > previous_change * 1.2:
                omega = max(0.125, omega * 0.5)
            if omega < 1.0:
                damped += 1
                nextT[:] = guessT + omega * (nextT - guessT)
                nextC[:] = guessC + omega * (nextC - guessC)
            if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                break
            if change <= 1.0:
                fill_properties(nextT, nextC, 2, a, k, D)
                links(k, h, faces, centers, gt)
                links(D, hm, faces, centers, gc)
                rt = residual(T, nextT, a, volume, gt, extT, dt, atol, rtol)
                rc = residual(C, nextC, ones, volume, gc, extC, dt, atol, rtol)
                if max(rt, rc) <= 1.0:
                    converged = True
                    worst = max(worst, rt, rc)
                    break
            guessT[:], guessC[:] = (nextT, nextC)
            previous_change = change
        iterations += it + 1
        if not converged:
            rejected += 1
            dt_limit = dt * 0.5
            if dt_limit < 1e-08:
                raise RuntimeError('Q3 Picard failed below minimum step; no clipping')
            continue
        loss += dt * gc[-1] * (nextC[-1] - extC) / total_volume
        T[:], C[:] = (nextT, nextC)
        t += dt
        minimum, maximum = (min(minimum, dt), max(maximum, dt))
        steps += 1
    diag = np.array([np.sum(volume * C) / total_volume, loss, float(steps), float(iterations), float(rejected), minimum if steps else 0.0, maximum, worst, float(damped)])
    return (T, C, diag)

@njit(cache=True)
def sample_fields(T, C, t, env, faces, centers, sample_r, h, hm):
    """Reconstruct output positions; full-grid states remain the solver state."""
    a, k, D = (np.empty(len(T)), np.empty(len(T)), np.empty(len(T)))
    fill_properties(T, C, 2, a, k, D)
    extT, extC = environment_at(t, env)
    return (reconstruct(T, k, h, extT, faces, centers, sample_r), reconstruct(C, D, hm, extC, faces, centers, sample_r))

@njit(cache=True)
def global_maximum(T, C, t, env, faces, centers, h, hm):
    """Maximum over all cells plus the symmetric center and Robin surface."""
    bounds = np.array([0.0, faces[-1]])
    _, boundary_C = sample_fields(T, C, t, env, faces, centers, bounds, h, hm)
    maximum, location = (boundary_C[0], 0.0)
    for j in range(len(C)):
        if C[j] > maximum:
            maximum, location = (C[j], centers[j])
    if boundary_C[1] > maximum:
        maximum, location = (boundary_C[1], faces[-1])
    return (maximum, location, boundary_C[0], boundary_C[1])

def refine_endpoint(T, C, t_start, t_stop, environment, faces, centers, volumes, config):
    """Evaluate every candidate from the same last-undried state, without interpolation.

    The returned diagnostic/loss belongs only to the accepted upper trajectory;
    work spent on discarded bisection candidates is reported separately.
    A coarse crossing can disappear under smaller steps. In that case `crossed`
    is false and the returned state is the re-integrated nonterminal upper state.
    """
    threshold = config['threshold']
    h, hm = (config['h'], config['hm'])
    m0, r0, cc0, cs0 = global_maximum(T, C, t_start, environment, faces, centers, h, hm)
    g0 = float(m0 - threshold)
    if g0 < 0.0:
        raise ValueError('Endpoint refinement requires an undried anchor, g >= 0')
    if not t_stop > t_start:
        raise ValueError('Endpoint interval must have positive length')
    trace = [[float(t_start), float(m0), float(r0), g0]]
    work_steps = 0
    work_iterations = 0

    def evaluate(t):
        nonlocal work_steps, work_iterations
        tn, cn, diag = advance_interval(T, C, float(t_start), float(t), config['event_dt'], environment, faces, centers, volumes, h, hm, config['atol'], config['rtol'], config['max_iterations'])
        m, r, cc, cs = global_maximum(tn, cn, t, environment, faces, centers, h, hm)
        trace.append([float(t), float(m), float(r), float(m - threshold)])
        work_steps += int(diag[2])
        work_iterations += int(diag[3])
        return {'T': tn, 'C': cn, 'diag': diag, 'M': float(m), 'r': float(r), 'C_center': float(cc), 'C_surface': float(cs), 'time': float(t)}
    upper = evaluate(t_stop)
    lower_T, lower_C = (T.copy(), C.copy())
    lower_diag = np.zeros(9)
    lo, hi = (float(t_start), float(t_stop))
    glo, ghi = (g0, upper['M'] - threshold)
    if ghi >= 0.0:
        return {**upper, 'crossed': False, 'trace': np.asarray(trace), 'bracket': None, 'work_steps': work_steps, 'work_iterations': work_iterations}
    while hi - lo > config['epsilon_t']:
        mid = lo + (hi - lo) * 0.5
        if mid == lo or mid == hi:
            raise RuntimeError('Endpoint tolerance is below floating-point time resolution')
        candidate = evaluate(mid)
        g = candidate['M'] - threshold
        if g < 0.0:
            hi, ghi, upper = (mid, g, candidate)
        else:
            lo, glo = (mid, g)
            lower_T, lower_C, lower_diag = (candidate['T'], candidate['C'], candidate['diag'])
    if not (glo >= 0.0 and ghi < 0.0):
        raise RuntimeError('A strict endpoint sign bracket was not established')
    return {**upper, 'crossed': True, 'trace': np.asarray(trace), 'bracket': {'t_minus': lo, 't_plus': hi, 'g_minus': glo, 'g_plus': ghi, 'width': hi - lo}, 'lower_T': lower_T, 'lower_C': lower_C, 'lower_diag': lower_diag, 'work_steps': work_steps, 'work_iterations': work_iterations}


def main(argv=None):
    parser=argparse.ArgumentParser(description="Standalone Q3 same-algorithm reproduction")
    folder=Path(__file__).resolve().parent
    parser.add_argument("--prefix",type=Path,default=folder.parent/"q2/solution.npz")
    parser.add_argument("--from-initial",action="store_true",
                        help="Recompute the full Q2 prefix using the standalone appendix kernel")
    parser.add_argument("--attachment",type=Path,default=folder.parent/"data/attachment1.xlsx")
    parser.add_argument("--template",type=Path,default=folder.parent/"data/templates/result3_template.xlsx")
    parser.add_argument("--output",type=Path,default=folder/"reproduced/q3")
    parser.add_argument("--dt-max",type=float,default=0.03125)
    parser.add_argument("--event-dt",type=float,default=0.05)
    parser.add_argument("--event-tol",type=float,default=0.01)
    args=parser.parse_args(argv)
    if min(args.dt_max,args.event_dt,args.event_tol)<=0:
        raise ValueError("Numerical scales must be positive")
    args.output.mkdir(parents=True,exist_ok=True)
    environment=load_environment(args.attachment)
    config=dict(question=2,end_time=10800,radius=.02,h=25.,hm=8e-7,
                atol=1e-10,rtol=1e-11,max_iterations=60,initial_T=301.15,
                initial_C=2.55,time_growth=600.)
    if args.from_initial:
        prefix=base.solve(config,environment,960,1/768,sample_count=21,progress=True)
        prefix_path=args.output/"q2_prefix_independent.npz"
        base.save_solution(prefix,prefix_path)
    else:
        prefix_path=args.prefix
        prefix=base.load_solution(args.prefix)
        assert prefix["metadata"]["question"]==2 and prefix["times"][-1]==10800.
        assert prefix["metadata"]["cells"]==960
        for key in ["h","hm","initial_T","initial_C"]:
            assert prefix["metadata"][key]==config[key]
    config.update(threshold=.15,event_dt=min(args.event_dt,args.dt_max),
                  epsilon_t=args.event_tol)
    faces,centers,V=prefix["faces"],prefix["centers"],prefix["volumes"]
    radii=np.linspace(0,.02,21)
    stride=(len(prefix["radii"])-1)//20
    assert np.allclose(prefix["radii"][::stride],radii,rtol=0,atol=1e-15)
    times=list(prefix["times"][::60])
    fieldsT=list(prefix["temperature_K"][::60,::stride])
    fieldsC=list(prefix["moisture"][::60,::stride])
    maxima=[float("nan")]*len(times)
    argmax=[float("nan")]*len(times)
    T,C=prefix["final_T"].copy(),prefix["final_C"].copy()
    t=10800.
    m,r,_,_=global_maximum(T,C,t,environment,faces,centers,25.,8e-7)
    maxima[-1],argmax[-1]=m,r
    loss=float(prefix["logs"][:,1].sum())
    prefix_cumulative=np.cumsum(prefix["logs"][:,1])
    balances=[0.]+list(prefix["logs"][59::60,0]+prefix_cumulative[59::60]-2.55)
    trace=[]
    last_print=time.perf_counter()
    while t<30*24*3600:
        next_time=t+60.
        tn,cn,diag=advance_interval(T,C,t,next_time,args.dt_max,environment,
                                   faces,centers,V,25.,8e-7,1e-10,1e-11,60)
        m,r,_,_=global_maximum(tn,cn,next_time,environment,faces,centers,25.,8e-7)
        terminal=False
        if m<.15:
            event=refine_endpoint(T,C,t,next_time,environment,faces,centers,V,config)
            trace.extend(event["trace"].tolist())
            tn,cn,diag=event["T"],event["C"],event["diag"]
            next_time,m,r=event["time"],event["M"],event["r"]
            terminal=event["crossed"]
        T,C,t=tn,cn,next_time
        loss+=float(diag[1])
        st,sc=sample_fields(T,C,t,environment,faces,centers,radii,25.,8e-7)
        times.append(t);fieldsT.append(st);fieldsC.append(sc)
        maxima.append(m);argmax.append(r)
        balances.append(float(np.dot(V,C)/V.sum()+loss-2.55))
        if terminal:
            bracket=event["bracket"]
            break
        if time.perf_counter()-last_print>30:
            print(f"Q3 appendix: t={t/3600:.3f} h, max C={m:.6f}",flush=True)
            last_print=time.perf_counter()
    else:
        raise RuntimeError("No endpoint by the computational safety bound; no success reported")
    assert bracket["g_minus"]>=0 and bracket["g_plus"]<0
    meta={"finish_time_s":t,"finish_time_h":t/3600,"bracket":bracket,
          "dt_max":args.dt_max,"event_dt":config["event_dt"],"event_tol":args.event_tol,
          "from_initial":args.from_initial,"standalone_same_algorithm":True,
          "prefix_path":str(prefix_path.resolve()),
          "prefix_sha256":hashlib.sha256(prefix_path.read_bytes()).hexdigest(),
          "source_attachment_sha256":hashlib.sha256(args.attachment.read_bytes()).hexdigest(),
          "max_water_balance":float(np.max(np.abs(balances)))}
    np.savez_compressed(args.output/"q3_independent.npz",
        times=np.asarray(times),radii=radii,temperature_K=np.asarray(fieldsT),
        moisture=np.asarray(fieldsC),max_C=np.asarray(maxima),max_r=np.asarray(argmax),
        final_T=T,final_C=C,faces=faces,centers=centers,volumes=V,
        water_balance=np.asarray(balances),event_trace=np.asarray(trace),
        metadata_json=json.dumps(meta,ensure_ascii=False))
    wb=load_workbook(args.template)
    ws=wb.worksheets[0]
    for row in ws:
        for cell in row:
            if cell.coordinate!="A1":cell.value=None
    for j in range(21):ws.cell(1,j+2,round(j*.1,1))
    output_row=2
    for ti,ci in zip(times,fieldsC):
        if ti<=0 or abs(ti/60-round(ti/60))>1e-9:continue
        ws.cell(output_row,1,int(round(ti)))
        for j,value in enumerate(ci):
            cell=ws.cell(output_row,j+2,round(float(value),4))
            cell.number_format="0.0000"
        output_row+=1
    ws.freeze_panes="B2"
    wb.properties.creator="";wb.properties.lastModifiedBy=""
    wb.save(args.output/"result3.xlsx")
    (args.output/"endpoint.json").write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(meta,ensure_ascii=False,indent=2),flush=True)
    return meta

if __name__=="__main__":
    main()
