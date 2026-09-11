"""Independent appendix solver: no imports from common, q1, q2 or run_all.
Includes precisely the same FVM, backward Euler, Picard and output reconstruction.
Dependencies: numpy, numba, openpyxl. Python 3.12 (64-bit) is the tested runtime.
"""
from pathlib import Path
import os,sys
APPENDIX_ROOT=Path(__file__).resolve().parent
runtime=APPENDIX_ROOT/".runtime"
for directory in [runtime/"tmp",runtime/"numba"]:directory.mkdir(parents=True,exist_ok=True)
os.environ["TEMP"]=os.environ["TMP"]=str(runtime/"tmp")
os.environ["NUMBA_CACHE_DIR"]=str(runtime/"numba")
sys.dont_write_bytecode=True
optional_deps=APPENDIX_ROOT.parent/".runtime/deps"
if optional_deps.exists():sys.path.insert(0,str(optional_deps))


# data_loader.py
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

def load_environment(path):
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    rows = list(wb.worksheets[0].values)
    wb.close()
    if tuple(rows[0]) != ("时间", "温度", "水分浓度"):
        raise ValueError("附件1表头不匹配")
    data = np.array(rows[1:], dtype=float)
    if data.shape[1] != 3 or not np.isfinite(data).all():
        raise ValueError("环境数据含空值或非有限数")
    if data[0, 0] != 0 or not np.all(np.diff(data[:, 0]) > 0):
        raise ValueError("环境时间必须从0开始且严格递增")
    data[:, 1] += 273.15
    return np.ascontiguousarray(data)


# interpolation.py
import numpy as np
from numba import njit

@njit(cache=True)
def environment_at(t, env):
    if t > env[-1, 0]:
        return 323.15, 0.0500
    return np.interp(t, env[:, 0], env[:, 1]), np.interp(t, env[:, 0], env[:, 2])


# material_properties.py
import numpy as np
from numba import njit

@njit(cache=True)
def properties(T, C, question):
    if question == 1:
        return 820.0 * 2600.0, 0.36, 7e-9 * np.exp(-0.89 / C)
    if question == 2:
        a = (650.0 + 128.0 * C) * (1450.0 + 2736.0 * C / (C + 1.0))
        k = 0.21 + 0.38 * C / (C + 1.0)
        D = 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / T)
        return a, k, D
    raise ValueError("Only questions 1 and 2 are implemented")


# boundary_conditions.py
from numba import njit

@njit(cache=True)
def surface_conductance(coefficient, exchange, distance, area):
    if exchange == 0.0:
        return 0.0
    return area / (distance / coefficient + 1.0 / exchange)

@njit(cache=True)
def surface_value(cell_value, environment, coefficient, exchange, distance):
    return cell_value + exchange * distance / (coefficient + exchange * distance) * (environment - cell_value)


# finite_volume.py
"""Conservative radial cell-centred FVM; backward Euler and Picard, no extrapolation."""
import numpy as np
from numba import njit

@njit(cache=True)
def geometry(n, radius):
    x = np.linspace(0.0, 1.0, n + 1)
    faces = radius * (10.0*x - x**10) / 9.0  # gentle interior spacing, refined surface layer
    volumes = (faces[1:]**2 - faces[:-1]**2) / 2.0  # omit common 2*pi*L
    centers = (2.0 / 3.0) * (faces[1:]**3 - faces[:-1]**3) / (faces[1:]**2 - faces[:-1]**2)
    return faces, centers, volumes

@njit(cache=True)
def fill_properties(T, C, question, a, k, D):
    for i in range(len(T)):
        a[i], k[i], D[i] = properties(T[i], C[i], question)

@njit(cache=True)
def links(coef, exchange, faces, centers, out):
    n = len(coef)
    out[0] = 0.0
    for j in range(1, n):
        out[j] = faces[j] / ((faces[j] - centers[j-1]) / coef[j-1] + (centers[j] - faces[j]) / coef[j])
    out[n] = surface_conductance(coef[-1], exchange, faces[-1]-centers[-1], faces[-1])

@njit(cache=True)
def linear_step(old, capacity, volume, conductance, external, dt, out, work):
    # Solve for the increment instead of absolute Kelvin values to avoid cancellation.
    n = len(old)
    upper, modified_rhs = work[0], work[1]
    for i in range(n):
        scale = dt / (capacity[i] * volume[i])
        left, right = conductance[i] * scale, conductance[i+1] * scale
        diagonal = 1.0 + left + right
        rhs = 0.0
        if i > 0:
            rhs += left * (old[i-1] - old[i])
        if i < n - 1:
            rhs += right * (old[i+1] - old[i])
        else:
            rhs += right * (external - old[i])
        if i > 0:
            diagonal -= left * upper[i-1]
            rhs += left * modified_rhs[i-1]
        upper[i] = right / diagonal
        modified_rhs[i] = rhs / diagonal
    increment = modified_rhs[-1]
    out[-1] = old[-1] + increment
    for i in range(n-2, -1, -1):
        increment = modified_rhs[i] + upper[i] * increment
        out[i] = old[i] + increment

@njit(cache=True)
def residual(old, new, capacity, volume, conductance, external, dt, atol, rtol):
    largest = 0.0
    for i in range(len(old)):
        mass = capacity[i] * volume[i] / dt
        flux = 0.0
        if i > 0:
            flux += conductance[i] * (new[i-1] - new[i])
        if i < len(old) - 1:
            flux += conductance[i+1] * (new[i+1] - new[i])
        else:
            flux += conductance[i+1] * (external - new[i])
        denom = (mass + conductance[i] + conductance[i+1]) * (atol + rtol*abs(new[i]))
        largest = max(largest, abs(mass*(new[i]-old[i])-flux)/denom)
    return largest

@njit(cache=True)
def reconstruct(values, coef, exchange, external, faces, centers, sample_r):
    # Even quadratic exactly matches the first two cylindrical volume averages.
    q0 = 0.5*(faces[0]**2 + faces[1]**2)
    q1 = 0.5*(faces[1]**2 + faces[2]**2)
    beta = (values[1]-values[0])/(q1-q0)
    center_value = values[0] - beta*q0
    surface = surface_value(values[-1], external, coef[-1], exchange, faces[-1]-centers[-1])
    xs = np.empty(len(values)+2)
    ys = np.empty(len(values)+2)
    xs[0], xs[-1] = 0.0, faces[-1]
    ys[0], ys[-1] = center_value, surface
    xs[1:-1], ys[1:-1] = centers, values
    result = np.interp(sample_r, xs, ys)
    for j in range(len(sample_r)):
        if sample_r[j] < centers[0]:
            result[j] = center_value + beta*sample_r[j]**2
    return result

@njit(cache=True)
def integrate_block(T, C, t_start, end_second, dt_max, question, env,
                    faces, centers, volume, sample_r, h, hm, atol, rtol, max_iterations, time_growth=0.0):
    n = len(T)
    count = end_second-int(t_start)
    ts = np.empty((count, len(sample_r)))
    cs = np.empty_like(ts)
    # columns: normalized mean C, normalized loss in this second, steps, iterations,
    # rejections, min dt, max dt, max scaled residual, damped iterations
    log = np.zeros((count, 9))
    a, k, D = np.empty(n), np.empty(n), np.empty(n)
    ones = np.ones(n)
    gt, gc = np.empty(n+1), np.empty(n+1)
    work = np.empty((2,n))
    guessT, guessC, nextT, nextC = np.empty(n), np.empty(n), np.empty(n), np.empty(n)
    total_volume = np.sum(volume)
    t = t_start
    for row in range(count):
        target = t_start+row+1
        loss, steps, iterations, rejected, minimum, maximum, worst, damped = 0.0,0,0,0,1e100,0.0,0.0,0
        dt_limit = min(1.0,4.0*dt_max) if time_growth > 0.0 else dt_max
        while t < target-1e-12:
            requested_dt = dt_max * (1.0+t/time_growth)**2 if time_growth > 0.0 else dt_max
            dt = min(dt_limit, requested_dt, target-t)
            # Attachment knots lie at whole seconds; integer output fences also respect them.
            extT, extC = environment_at(t+dt,env)
            guessT[:], guessC[:] = T, C
            converged = False
            previous_change = 1e100
            omega = 1.0
            for it in range(max_iterations):
                fill_properties(guessT,guessC,question,a,k,D)
                if question == 2 or it == 0:
                    links(k,h,faces,centers,gt)
                    linear_step(T,a,volume,gt,extT,dt,nextT,work)
                fill_properties(nextT,guessC,question,a,k,D)
                links(D,hm,faces,centers,gc)
                linear_step(C,ones,volume,gc,extC,dt,nextC,work)
                if not np.all(np.isfinite(nextT)) or not np.all(np.isfinite(nextC)):
                    break
                if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                    omega *= 0.5
                change = 0.0
                for j in range(n):
                    change=max(change,abs(nextT[j]-guessT[j])/(atol+rtol*abs(nextT[j])))
                    change=max(change,abs(nextC[j]-guessC[j])/(atol+rtol*abs(nextC[j])))
                if change > previous_change*1.2:
                    omega = max(0.125,omega*0.5)
                if omega < 1.0:
                    damped += 1
                    nextT[:] = guessT + omega*(nextT-guessT)
                    nextC[:] = guessC + omega*(nextC-guessC)
                if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                    break
                if change <= 1.0:
                    fill_properties(nextT,nextC,question,a,k,D)
                    links(k,h,faces,centers,gt)
                    links(D,hm,faces,centers,gc)
                    rt = residual(T,nextT,a,volume,gt,extT,dt,atol,rtol)
                    rc = residual(C,nextC,ones,volume,gc,extC,dt,atol,rtol)
                    if max(rt,rc) <= 1.0:
                        converged=True
                        worst=max(worst,rt,rc)
                        break
                guessT[:],guessC[:] = nextT,nextC
                previous_change=change
            iterations += it+1
            if not converged:
                rejected += 1
                dt_limit=dt*0.5
                if dt_limit < 1e-8:
                    raise RuntimeError("Picard failed below minimum time step; no state clipping applied")
                continue
            loss += dt*gc[-1]*(nextC[-1]-extC)/total_volume
            T[:],C[:] = nextT,nextC
            t += dt
            minimum,maximum=min(minimum,dt),max(maximum,dt)
            steps += 1
        t=target
        fill_properties(T,C,question,a,k,D)
        extT,extC=environment_at(t,env)
        ts[row]=reconstruct(T,k,h,extT,faces,centers,sample_r)
        cs[row]=reconstruct(C,D,hm,extC,faces,centers,sample_r)
        log[row,0]=np.sum(volume*C)/total_volume
        log[row,1]=loss
        log[row,2]=steps
        log[row,3]=iterations
        log[row,4]=rejected
        log[row,5]=minimum
        log[row,6]=maximum
        log[row,7]=worst
        log[row,8]=damped
    return T,C,ts,cs,log


# solver.py
from pathlib import Path
import json, time
import numpy as np

def solve(config, environment, cells, dt, sample_count=21, progress=True):
    started=time.perf_counter()
    faces,centers,volumes=geometry(cells,config["radius"])
    radii=np.linspace(0,config["radius"],sample_count)
    end=int(config["end_time"])
    T=np.full(cells,config["initial_T"])
    C=np.full(cells,config["initial_C"])
    outputT=np.empty((end+1,sample_count));outputC=np.empty_like(outputT)
    outputT[0]=config["initial_T"];outputC[0]=config["initial_C"]
    logs=np.zeros((end,9))
    last_report=time.perf_counter()
    for start in range(0,end,60):
        stop=min(start+60,end)
        T,C,ot,oc,log=integrate_block(T,C,float(start),stop,dt,config["question"],
            environment,faces,centers,volumes,radii,config["h"],config["hm"],
            config["atol"],config["rtol"],config["max_iterations"],config.get("time_growth",0.0))
        outputT[start+1:stop+1]=ot;outputC[start+1:stop+1]=oc;logs[start:stop]=log
        if progress and time.perf_counter()-last_report>30:
            print(f'q{config["question"]} N={cells} dt={dt:g}: {stop}/{end} s',flush=True)
            last_report=time.perf_counter()
    cumulative=np.cumsum(logs[:,1])
    balance=logs[:,0]+cumulative-config["initial_C"]
    meta={**config,"mesh":"polynomial10_surface_graded","cells":cells,"dt_scale":dt,"dt_max":min(1.0,4.0*dt,dt*(1+end/config["time_growth"])**2) if config.get("time_growth",0)>0 else dt,"sample_count":sample_count,
          "elapsed_seconds":time.perf_counter()-started,"steps":int(logs[:,2].sum()),
          "iterations":int(logs[:,3].sum()),"rejections":int(logs[:,4].sum()),
          "damped_iterations":int(logs[:,8].sum()),
          "max_scaled_residual":float(logs[:,7].max()),
          "max_water_balance_residual":float(np.max(np.abs(balance)))}
    return {"times":np.arange(end+1,dtype=float),"radii":radii,"temperature_K":outputT,
            "moisture":outputC,"faces":faces,"centers":centers,"volumes":volumes,
            "final_T":T,"final_C":C,"logs":logs,"water_balance":balance,"metadata":meta}

def save_solution(result,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    np.savez_compressed(path,**{k:v for k,v in result.items() if k!="metadata"},
                        metadata_json=json.dumps(result["metadata"],ensure_ascii=False))
    path.with_suffix(".json").write_text(json.dumps(result["metadata"],ensure_ascii=False,indent=2),encoding="utf-8")

def load_solution(path):
    with np.load(path,allow_pickle=False) as data:
        result={k:data[k] for k in data.files if k!="metadata_json"}
        result["metadata"]=json.loads(str(data["metadata_json"]))
        return result


# excel_export.py
"""Template-preserving numerical outputs; no solver calculations in Excel."""
from pathlib import Path
from copy import copy
from openpyxl import load_workbook

def export_result(solution,template_path,destination):
    wb=load_workbook(template_path)
    for sheet,key in [("温度","temperature_K"),("水分浓度","moisture")]:
        ws=wb[sheet]
        # Replace illustrative ellipses; source template file is never edited.
        for row in ws:
            for cell in row:
                if cell.coordinate != "A1": cell.value=None
        for j in range(21):ws.cell(1,j+2,round(j*.1,1))
        stride=(len(solution["radii"])-1)//20
        values=solution[key][:,::stride]
        for i in range(1,len(solution["times"])):
            ws.cell(i+1,1,int(solution["times"][i]))
            for j in range(21):
                value=float(values[i,j])-(273.15 if key=="temperature_K" else 0)
                cell=ws.cell(i+1,j+2,round(value,4));cell.number_format="0.0000"
        ws.freeze_panes="B2"
        ws.column_dimensions["A"].width=28
        from openpyxl.utils import get_column_letter
        for j in range(2,23):ws.column_dimensions[get_column_letter(j)].width=11
        ws.row_dimensions[1].height=32
        from openpyxl.styles import Alignment
        ws["A1"].alignment=Alignment(wrap_text=True,vertical="center")
    wb.properties.creator="";wb.properties.lastModifiedBy=""
    destination=Path(destination);destination.parent.mkdir(parents=True,exist_ok=True)
    wb.save(destination)


def main(argv=None):
    import argparse
    parser=argparse.ArgumentParser(description="独立附录代码：第一、二问完整结果")
    parser.add_argument("--question",choices=["1","2","all"],default="all")
    parser.add_argument("--attachment",type=Path,default=APPENDIX_ROOT.parent/"data/attachment1.xlsx")
    parser.add_argument("--templates",type=Path,default=APPENDIX_ROOT.parent/"data/templates")
    parser.add_argument("--output",type=Path,default=APPENDIX_ROOT/"reproduced")
    parser.add_argument("--compute-only",action="store_true")
    args=parser.parse_args(argv)
    env=load_environment(args.attachment)
    for q in ([1,2] if args.question=="all" else [int(args.question)]):
        config={"question":q,"end_time":1800 if q==1 else 10800,"radius":.02,
                "h":25.,"hm":8e-7,"atol":1e-10,"rtol":1e-11,"max_iterations":60,
                "initial_T":301.15,"initial_C":2.55,"time_growth":600.}
        result=solve(config,env,960,1/768,sample_count=21)
        args.output.mkdir(parents=True,exist_ok=True)
        save_solution(result,args.output/f"q{q}_independent.npz")
        if not args.compute_only:
            export_result(result,args.templates/f"result{q}_template.xlsx",args.output/f"result{q}.xlsx")
        print(f"Independent question {q} complete",flush=True)
if __name__=="__main__":main()
