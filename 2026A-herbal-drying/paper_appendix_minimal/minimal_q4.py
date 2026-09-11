"""Standalone Q4 material-domain solver; see the appendix README."""
from pathlib import Path
import argparse, json, time, hashlib
import minimal_solver as base
import numpy as np
from numba import njit
from openpyxl import load_workbook
from minimal_solver import geometry, links, linear_step, residual, reconstruct, environment_at, load_environment

SELECTED = {'cells': 1280, 'dt_max': 0.03125, 'dt_scale': 0.0013020833333333333, 'event_dt': 0.03125, 'event_tol': 0.01}

def pchip_slopes(data):
    x, y = (data[:, 0], data[:, 1])
    h = np.diff(x)
    delta = np.diff(y) / h
    d = np.zeros(len(x))
    if len(x) == 2:
        d[:] = delta[0]
        return d
    for i in range(1, len(x) - 1):
        if delta[i - 1] * delta[i] > 0:
            w1, w2 = (2 * h[i] + h[i - 1], h[i] + 2 * h[i - 1])
            d[i] = (w1 + w2) / (w1 / delta[i - 1] + w2 / delta[i])
    for i, j, k in ((0, 0, 1), (-1, -1, -2)):
        d[i] = ((2 * h[j] + h[k]) * delta[j] - h[j] * delta[k]) / (h[j] + h[k])
        if np.sign(d[i]) != np.sign(delta[j]):
            d[i] = 0.0
        elif np.sign(delta[j]) != np.sign(delta[k]) and abs(d[i]) > 3 * abs(delta[j]):
            d[i] = 3 * delta[j]
    return d

def load_radius_data(path):
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    rows = [(float(t), float(r) * 0.01) for t, r, *_ in wb.active.iter_rows(min_row=2, values_only=True) if t is not None]
    wb.close()
    data = np.asarray(rows, dtype=np.float64)
    if data.ndim != 2 or data.shape[1] != 2 or len(data) < 2 or (not np.all(np.isfinite(data))) or (data[0, 0] != 0) or np.any(np.diff(data[:, 0]) <= 0) or np.any(data[:, 1] <= 0) or np.any(np.diff(data[:, 1]) > 0):
        raise ValueError('Radius input requires increasing seconds and positive nonincreasing metres')
    return (data, pchip_slopes(data))

@njit(cache=True)
def radius_at(t, data, slopes, shrinking=True, radius0=0.02):
    if not shrinking:
        return radius0
    if t <= data[0, 0]:
        return data[0, 1]
    if t >= data[-1, 0]:
        return data[-1, 1]
    j = np.searchsorted(data[:, 0], t, side='right') - 1
    if data[j, 1] == data[j + 1, 1]:
        return data[j, 1]
    h = data[j + 1, 0] - data[j, 0]
    z = (t - data[j, 0]) / h
    return (2 * z ** 3 - 3 * z * z + 1) * data[j, 1] + (z ** 3 - 2 * z * z + z) * h * slopes[j] + (-2 * z ** 3 + 3 * z * z) * data[j + 1, 1] + (z ** 3 - z * z) * h * slopes[j + 1]

@njit(cache=True)
def fill_properties4(T, C, scale, a, k, D):
    """Appendix 4; scale transports once on the fixed reference domain."""
    for j in range(len(T)):
        c = C[j]
        fraction = c / (1.0 + c)
        a[j] = (760.0 + 90.0 * c) * (1850.0 + 2150.0 * fraction)
        k[j] = (0.12 + 0.2 * fraction) / (scale * scale)
        D[j] = 0.00042 * np.exp(-0.3 / c) * np.exp(-3850.0 / T[j]) / (scale * scale)

@njit(cache=True)
def advance_interval(T0, C0, t_start, t_end, dt_max, env, faces, centers, volume, h, hm, atol, rtol, max_iterations, radius_data, radius_slopes, shrinking=True, dt_scale=1.0 / 768.0, time_growth=600.0):
    """Return new full-grid states and 9 diagnostics; never mutate input state.

    Diagnostic order agrees with integrate_block:
    mean_C, interval_loss, steps, iterations, rejections, min_dt, max_dt,
    worst_scaled_residual, damped_iterations.
    """
    if t_end < t_start or dt_max <= 0.0 or max_iterations < 1:
        raise ValueError('Invalid continuation interval or step configuration')
    T, C = (T0.copy(), C0.copy())
    n = len(T)
    a, k, D = (np.empty(n), np.empty(n), np.empty(n))
    ones = np.ones(n)
    gt, gc = (np.empty(n + 1), np.empty(n + 1))
    work = np.empty((2, n))
    guessT, guessC = (np.empty(n), np.empty(n))
    nextT, nextC = (np.empty(n), np.empty(n))
    total_volume = np.sum(volume)
    t, dt_limit = (t_start, dt_max)
    loss, steps, iterations, rejected = (0.0, 0, 0, 0)
    minimum, maximum, worst, damped = (1e+100, 0.0, 0.0, 0)
    while t < t_end - 1e-12:
        knot_index = np.searchsorted(env[:, 0], t + 1e-10, side='right')
        next_knot = env[knot_index, 0] if knot_index < len(env) else t_end
        radius_index = np.searchsorted(radius_data[:, 0], t + 1e-10, side='right')
        radius_knot = radius_data[radius_index, 0] if radius_index < len(radius_data) else t_end
        requested = dt_max if time_growth <= 0 else dt_scale * (1.0 + t / time_growth) ** 2
        dt = min(dt_limit, requested, t_end - t, next_knot - t, radius_knot - t)
        scale = radius_at(t + dt, radius_data, radius_slopes, shrinking, faces[-1]) / faces[-1]
        href, hmref = (h / scale, hm / scale)
        extT, extC = environment_at(t + dt, env)
        guessT[:], guessC[:] = (T, C)
        converged, previous_change, omega = (False, 1e+100, 1.0)
        for it in range(max_iterations):
            fill_properties4(guessT, guessC, scale, a, k, D)
            links(k, href, faces, centers, gt)
            linear_step(T, a, volume, gt, extT, dt, nextT, work)
            fill_properties4(nextT, guessC, scale, a, k, D)
            links(D, hmref, faces, centers, gc)
            linear_step(C, ones, volume, gc, extC, dt, nextC, work)
            if not np.all(np.isfinite(nextT)) or not np.all(np.isfinite(nextC)):
                break
            if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                omega *= 0.5
            change = 0.0
            for j in range(n):
                change = max(change, abs(nextT[j] - guessT[j]) / (atol + rtol * abs(nextT[j])))
                change = max(change, abs(nextC[j] - guessC[j]) / (atol + rtol * abs(nextC[j])))
            if change > previous_change * 1.2:
                omega = max(0.125, omega * 0.5)
            if omega < 1.0:
                damped += 1
                nextT[:] = guessT + omega * (nextT - guessT)
                nextC[:] = guessC + omega * (nextC - guessC)
            if np.min(nextC) <= 0.0 or np.min(nextT) <= 0.0:
                break
            if change <= 1.0:
                fill_properties4(nextT, nextC, scale, a, k, D)
                links(k, href, faces, centers, gt)
                links(D, hmref, faces, centers, gc)
                rt = residual(T, nextT, a, volume, gt, extT, dt, atol, rtol)
                rc = residual(C, nextC, ones, volume, gc, extC, dt, atol, rtol)
                if max(rt, rc) <= 1.0:
                    converged = True
                    worst = max(worst, rt, rc)
                    break
            guessT[:], guessC[:] = (nextT, nextC)
            previous_change = change
        iterations += it + 1
        if not converged:
            rejected += 1
            dt_limit = dt * 0.5
            if dt_limit < 1e-08:
                raise RuntimeError('Q4 Picard failed below minimum step; no clipping')
            continue
        loss += dt * gc[-1] * (nextC[-1] - extC) / total_volume
        T[:], C[:] = (nextT, nextC)
        t += dt
        minimum, maximum = (min(minimum, dt), max(maximum, dt))
        steps += 1
    diag = np.array([np.sum(volume * C) / total_volume, loss, float(steps), float(iterations), float(rejected), minimum if steps else 0.0, maximum, worst, float(damped)])
    return (T, C, diag)

@njit(cache=True)
def sample_fields(T, C, t, env, faces, centers, sample_r, h, hm, radius_data, radius_slopes, shrinking=True, reference=False):
    """Mask physical exterior BEFORE reconstruction; reference points are metres."""
    R = radius_at(t, radius_data, radius_slopes, shrinking, faces[-1])
    scale = R / faces[-1]
    mask = (sample_r >= 0.0) & (sample_r <= (faces[-1] if reference else R))
    indices = np.nonzero(mask)[0]
    query = sample_r[indices].copy()
    if not reference:
        query /= scale
    outT, outC = (np.full(len(sample_r), np.nan), np.full(len(sample_r), np.nan))
    if t == 0.0 and np.max(T) == np.min(T) and (np.max(C) == np.min(C)):
        outT[indices], outC[indices] = (T[0], C[0])
        return (outT, outC, mask)
    a, k, D = (np.empty(len(T)), np.empty(len(T)), np.empty(len(T)))
    fill_properties4(T, C, scale, a, k, D)
    extT, extC = environment_at(t, env)
    outT[indices] = reconstruct(T, k, h / scale, extT, faces, centers, query)
    outC[indices] = reconstruct(C, D, hm / scale, extC, faces, centers, query)
    return (outT, outC, mask)

@njit(cache=True)
def surface_fields(T, C, t, env, faces, centers, h, hm, radius_data, radius_slopes, shrinking=True):
    ts, cs, _ = sample_fields(T, C, t, env, faces, centers, np.array([faces[-1]]), h, hm, radius_data, radius_slopes, shrinking, True)
    return (ts[0], cs[0])

@njit(cache=True)
def global_maximum(T, C, t, env, faces, centers, h, hm, radius_data, radius_slopes, shrinking=True):
    _, bounds, _ = sample_fields(T, C, t, env, faces, centers, np.array([0.0, faces[-1]]), h, hm, radius_data, radius_slopes, shrinking, True)
    maximum, location = (bounds[0], 0.0)
    for j in range(len(C)):
        if C[j] > maximum:
            maximum, location = (C[j], centers[j])
    if bounds[1] > maximum:
        maximum, location = (bounds[1], faces[-1])
    scale = radius_at(t, radius_data, radius_slopes, shrinking, faces[-1]) / faces[-1]
    return (maximum, scale * location, bounds[0], bounds[1])

def refine_endpoint(T, C, t_start, t_stop, environment, faces, centers, volumes, config, radius_data, radius_slopes):
    """Recompute each candidate from the same full-grid last-undried state.

    No state or crossing time is interpolated. A coarse crossing that vanishes
    under the event step returns a nonterminal re-integrated upper state.
    Only the accepted trajectory contributes to cumulative water loss.
    """
    h, hm = (config['h'], config['hm'])
    threshold, shrinking = (config['threshold'], config['shrinking'])
    m0, r0, _, _ = global_maximum(T, C, t_start, environment, faces, centers, h, hm, radius_data, radius_slopes, shrinking)
    if not np.isfinite(m0) or m0 < threshold:
        raise ValueError('Endpoint refinement needs a finite undried anchor with g >= 0')
    if not t_stop > t_start:
        raise ValueError('Endpoint interval must have positive length')
    trace = [[float(t_start), float(m0), float(r0), float(m0 - threshold)]]
    work_steps = work_iterations = 0

    def evaluate(t):
        nonlocal work_steps, work_iterations
        tn, cn, diag = advance_interval(T, C, float(t_start), float(t), config['event_dt'], environment, faces, centers, volumes, h, hm, config['atol'], config['rtol'], config['max_iterations'], radius_data, radius_slopes, shrinking, config['dt_scale'], config['time_growth'])
        m, r, cc, cs = global_maximum(tn, cn, t, environment, faces, centers, h, hm, radius_data, radius_slopes, shrinking)
        if not np.isfinite(m) or not np.isfinite(diag).all():
            raise RuntimeError('Nonfinite event state or diagnostics')
        trace.append([float(t), float(m), float(r), float(m - threshold)])
        work_steps += int(diag[2])
        work_iterations += int(diag[3])
        return {'T': tn, 'C': cn, 'diag': diag, 'M': float(m), 'r': float(r), 'C_center': float(cc), 'C_surface': float(cs), 'time': float(t)}
    upper = evaluate(t_stop)
    lo, hi = (float(t_start), float(t_stop))
    glo, ghi = (float(m0 - threshold), upper['M'] - threshold)
    lower_T, lower_C, lower_diag = (T.copy(), C.copy(), np.zeros(9))
    if ghi >= 0.0:
        return {**upper, 'crossed': False, 'trace': np.asarray(trace), 'bracket': None, 'work_steps': work_steps, 'work_iterations': work_iterations}
    while hi - lo > config['epsilon_t']:
        mid = lo + (hi - lo) * 0.5
        if mid == lo or mid == hi:
            raise RuntimeError('Endpoint tolerance is below floating-point time resolution')
        candidate = evaluate(mid)
        g = candidate['M'] - threshold
        if g < 0.0:
            hi, ghi, upper = (mid, g, candidate)
        else:
            lo, glo = (mid, g)
            lower_T, lower_C, lower_diag = (candidate['T'], candidate['C'], candidate['diag'])
    if not (glo >= 0.0 and ghi < 0.0):
        raise RuntimeError('A strict endpoint sign bracket was not established')
    return {**upper, 'crossed': True, 'trace': np.asarray(trace), 'bracket': {'t_minus': lo, 't_plus': hi, 'g_minus': glo, 'g_plus': ghi, 'width': hi - lo}, 'lower_T': lower_T, 'lower_C': lower_C, 'lower_diag': lower_diag, 'work_steps': work_steps, 'work_iterations': work_iterations}

def main(argv=None):
    folder=Path(__file__).resolve().parent
    p=argparse.ArgumentParser(description="Standalone appendix-4 reproduction from t=0")
    p.add_argument('--attachment',type=Path,default=folder.parent/'data/attachment1.xlsx')
    p.add_argument('--radius',type=Path,default=folder.parent/'data/attachment2.xlsx')
    p.add_argument('--template',type=Path,default=folder.parent/'data/templates/result4_template.xlsx')
    p.add_argument('--output',type=Path,default=folder/'reproduced/q4')
    p.add_argument('--cells',type=int,default=SELECTED['cells'])
    p.add_argument('--dt-max',type=float,default=SELECTED['dt_max'])
    p.add_argument('--dt-scale',type=float,default=SELECTED['dt_scale'])
    p.add_argument('--event-dt',type=float,default=SELECTED['event_dt'])
    p.add_argument('--event-tol',type=float,default=SELECTED['event_tol'])
    p.add_argument('--fixed-radius',action='store_true')
    args=p.parse_args(argv)
    if args.cells<2 or min(args.dt_max,args.dt_scale,args.event_dt,args.event_tol)<=0:
        raise ValueError('Positive steps and at least two cells are required')
    args.output.mkdir(parents=True,exist_ok=True)
    env=load_environment(args.attachment); rd,rs=load_radius_data(args.radius)
    faces,centers,V=geometry(args.cells,.02)
    T,C=np.full(args.cells,301.15),np.full(args.cells,2.55)
    radii=np.linspace(0,.02,21)
    cfg=dict(h=25.,hm=8e-7,threshold=.15,shrinking=not args.fixed_radius,
             atol=1e-10,rtol=1e-11,max_iterations=60,dt_scale=args.dt_scale,
             time_growth=600.,event_dt=min(args.event_dt,args.dt_max),epsilon_t=args.event_tol)
    times=[];Ts=[];Cs=[];surfaceT=[];surfaceC=[];rhistory=[];masks=[]
    maxima=[];locations=[];balances=[];losses=[]
    def record(t,loss):
        st,sc,mask=sample_fields(T,C,t,env,faces,centers,radii,25.,8e-7,rd,rs,cfg['shrinking'])
        ut,uc=surface_fields(T,C,t,env,faces,centers,25.,8e-7,rd,rs,cfg['shrinking'])
        m,r,_,_=global_maximum(T,C,t,env,faces,centers,25.,8e-7,rd,rs,cfg['shrinking'])
        times.append(t);Ts.append(st);Cs.append(sc);masks.append(mask)
        surfaceT.append(ut);surfaceC.append(uc);rhistory.append(radius_at(t,rd,rs,cfg['shrinking']))
        maxima.append(m);locations.append(r);losses.append(loss)
        balances.append(np.dot(V,C)/V.sum()+loss-2.55)
    t,loss=0.,0.;record(t,loss)
    last_print=time.perf_counter();started=last_print;end_limit=72*3600.
    trace=[]
    while t<30*86400.:
        if t>=end_limit:
            end_limit=min(2*end_limit,30*86400.)
            print(f'Extending search to {end_limit/3600:g} h',flush=True)
        stop=t+60.
        tn,cn,diag=advance_interval(T,C,t,stop,args.dt_max,env,faces,centers,V,
            25.,8e-7,1e-10,1e-11,60,rd,rs,cfg['shrinking'],args.dt_scale,600.)
        m,_,_,_=global_maximum(tn,cn,stop,env,faces,centers,25.,8e-7,rd,rs,cfg['shrinking'])
        terminal=False
        if m<.15:
            event=refine_endpoint(T,C,t,stop,env,faces,centers,V,cfg,rd,rs)
            trace.extend(event['trace'].tolist())
            tn,cn,diag,stop=event['T'],event['C'],event['diag'],event['time']
            terminal=event['crossed']
        T,C,t=tn,cn,stop;loss+=diag[1];record(t,loss)
        if terminal:break
        if time.perf_counter()-last_print>30:
            print(f'Q4 appendix: t={t/3600:.3f} h; max C={maxima[-1]:.8f}',flush=True)
            last_print=time.perf_counter()
    else:
        raise RuntimeError('No threshold crossing within computational search bound')
    bracket=event['bracket']
    assert bracket['g_minus']>=0 and bracket['g_plus']<0
    meta=dict(finish_time_s=t,finish_time_h=t/3600,bracket=bracket,cells=args.cells,
        dt_max=args.dt_max,dt_scale=args.dt_scale,event_dt=cfg['event_dt'],event_tol=args.event_tol,
        shrinking=cfg['shrinking'],from_initial=True,standalone_same_algorithm=True,
        maximum_water_balance=float(np.max(np.abs(balances))),runtime_s=time.perf_counter()-started,
        attachment_sha256=hashlib.sha256(args.attachment.read_bytes()).hexdigest(),
        radius_sha256=hashlib.sha256(args.radius.read_bytes()).hexdigest())
    np.savez_compressed(args.output/'q4_independent.npz',times=np.asarray(times),radii=radii,
        temperature_K=np.asarray(Ts),moisture=np.asarray(Cs),domain_mask=np.asarray(masks),
        surface_temperature_K=np.asarray(surfaceT),surface_moisture=np.asarray(surfaceC),
        radius_m=np.asarray(rhistory),max_C=np.asarray(maxima),max_r=np.asarray(locations),
        water_balance=np.asarray(balances),loss_cumulative=np.asarray(losses),
        final_T=T,final_C=C,faces=faces,centers=centers,volumes=V,
        radius_data=rd,radius_slopes=rs,event_trace=np.asarray(trace),metadata_json=json.dumps(meta))
    wb=load_workbook(args.template);ws=wb.active
    for row in ws:
        for cell in row:
            if cell.coordinate!='A1':cell.value=None
    for j in range(21):ws.cell(1,j+2,round(j*.1,1))
    ws.cell(1,23,'药材表面');row=2
    for index,ti in enumerate(times):
        if ti<=0 or abs(ti/60-round(ti/60))>1e-9:continue
        ws.cell(row,1,int(round(ti)))
        for j,value in enumerate(list(Cs[index])+[surfaceC[index]]):
            cell=ws.cell(row,j+2)
            if np.isfinite(value):cell.value=round(float(value),4)
            cell.number_format='0.0000'
        row+=1
    ws.freeze_panes='B2';wb.properties.creator='';wb.properties.lastModifiedBy=''
    wb.save(args.output/'result4.xlsx')
    (args.output/'endpoint.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2),flush=True)
    return meta

if __name__=='__main__':main()
