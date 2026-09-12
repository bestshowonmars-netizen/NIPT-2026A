"""Q3 paper tables, template workbook, diagnostics, and figure from saved states."""
from pathlib import Path
import csv
import json

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from q3.plotting import figure5


def _json_value(value):
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, np.ndarray):
        return _json_value(value.tolist())
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    return value


def _write_json(path, payload):
    Path(path).write_text(json.dumps(_json_value(payload), ensure_ascii=False,
                                    indent=2, allow_nan=False) + "\n", encoding="utf-8")


def _exact_indices(available, requested, atol, description):
    """Match existing states; never interpolate or use a rounded workbook."""
    available = np.asarray(available, dtype=float)
    requested = np.asarray(requested, dtype=float)
    indices = np.searchsorted(available, requested)
    indices = np.minimum(indices, len(available) - 1)
    previous = np.maximum(indices - 1, 0)
    use_previous = np.abs(available[previous] - requested) < np.abs(available[indices] - requested)
    indices[use_previous] = previous[use_previous]
    if not np.allclose(available[indices], requested, rtol=0, atol=atol):
        missed = requested[np.abs(available[indices] - requested) > atol]
        raise ValueError(f"Missing saved {description}: {missed.tolist()}")
    return indices


def _check_solution(solution):
    times = np.asarray(solution["times"], dtype=float)
    radii = np.asarray(solution["radii"], dtype=float)
    if times.ndim != 1 or not len(times) or not np.all(np.isfinite(times)):
        raise ValueError("Q3 times must be a finite one-dimensional array.")
    if times[0] != 0 or not np.all(np.diff(times) > 0):
        raise ValueError("Q3 output must contain the initial state and increasing absolute times.")
    if radii.ndim != 1 or not np.all(np.diff(radii) > 0):
        raise ValueError("Q3 radii must be increasing actual positions in metres.")
    if not np.isclose(radii[0], 0, atol=1e-12) or not np.isclose(radii[-1], .02, rtol=0, atol=1e-12):
        raise ValueError("The Q3 fixed-radius output must cover r=0 to 0.02 m.")
    for name in ("temperature_K", "moisture"):
        values = np.asarray(solution[name])
        if values.shape != (len(times), len(radii)) or not np.all(np.isfinite(values)):
            raise ValueError(f"Invalid saved field: {name}")
    for name in ("max_C", "max_r", "mean_C", "loss_cumulative", "water_balance"):
        if np.asarray(solution[name]).shape != times.shape:
            raise ValueError(f"Invalid Q3 history shape: {name}")
    if np.asarray(solution["logs"]).shape != (len(times), 9):
        raise ValueError("Q3 interval diagnostics must have nine fields per saved time.")
    meta = solution["metadata"]
    finish = float(meta["finish_time_s"])
    threshold = float(meta.get("threshold", .15))
    if not np.isclose(times[-1], finish, rtol=0, atol=1e-7):
        raise ValueError("Last saved time does not match the precise endpoint.")
    maximum = float(solution["max_C"][-1])
    maximum_r = float(solution["max_r"][-1])
    if not np.isfinite(maximum) or not maximum < threshold:
        raise ValueError("The unrounded final global maximum must be strictly below 0.15.")
    if not np.isfinite(maximum_r) or not radii[0] <= maximum_r <= radii[-1]:
        raise ValueError("The endpoint needs an actual global-maximum location.")
    bracket = meta["bracket"]
    if not (float(bracket["g_minus"]) >= 0 and float(bracket["g_plus"]) < 0):
        raise ValueError("The endpoint must have an undried/dried sign bracket.")
    if not np.isclose(float(bracket["t_plus"]), finish, rtol=0, atol=1e-7):
        raise ValueError("The reported finish must be the strictly dried bracket upper state.")
    return times, radii, finish, threshold


def export_result3(solution, template_path, destination):
    """Preserve Sheet1 and write only 60, 120, ... second states to result3.xlsx."""
    times, radii, finish, _ = _check_solution(solution)
    minute_times = np.arange(1, int(np.floor(finish / 60.0)) + 1, dtype=np.int64) * 60
    time_indices = _exact_indices(times, minute_times, 1e-7, "whole-minute times")
    radial_indices = _exact_indices(radii, np.arange(21) * .001, 1e-12, "workbook radii")
    values = np.asarray(solution["moisture"])[np.ix_(time_indices, radial_indices)]
    wb = load_workbook(template_path)
    if wb.sheetnames != ["Sheet1"]:
        raise ValueError("result3 template must contain only Sheet1.")
    ws = wb["Sheet1"]
    title = ws["A1"].value or "时间\\到药材中心的距离"
    # Clear illustrative ellipses and old rows, while preserving the supplied template.
    for row in ws:
        for cell in row:
            cell.value = None
    ws["A1"] = title
    for j in range(21):
        ws.cell(1, j + 2, round(j * .1, 1)).number_format = "0.0"
    for i, second in enumerate(minute_times, start=2):
        ws.cell(i, 1, int(second)).number_format = "0"
        for j, concentration in enumerate(values[i - 2], start=2):
            ws.cell(i, j, round(float(concentration), 4)).number_format = "0.0000"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 28
    ws.row_dimensions[1].height = 32
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    for column in range(2, 23):
        ws.column_dimensions[get_column_letter(column)].width = 11
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    wb.close()
    return {"path": str(destination), "sheet": "Sheet1", "data_rows": len(minute_times),
            "first_time_s": int(minute_times[0]) if len(minute_times) else None,
            "last_time_s": int(minute_times[-1]) if len(minute_times) else None,
            "radial_columns": 21, "nonminute_endpoint_included": False}


def _write_table5(solution, folder, times, radii, finish):
    positions = np.array([0, .005, .01, .015, .02])
    columns = _exact_indices(radii, positions, 1e-12, "table 5 radii")
    targets = np.arange(1, int(np.floor(finish / 21600.0)) + 1, dtype=float) * 21600
    # A terminal state that is exactly a 6-hour point appears once, as the endpoint.
    targets = targets[targets < finish - 1e-7]
    row_indices = list(_exact_indices(times, targets, 1e-7, "six-hour states")) + [len(times) - 1]
    header = ["time_s", "time_h", "row_kind"] + [f"r={r:g} cm" for r in positions * 100]
    rows = []
    for k, idx in enumerate(row_indices):
        kind = "endpoint" if k == len(row_indices) - 1 else "regular_6h"
        values = np.asarray(solution["moisture"])[idx, columns]
        rows.append([float(times[idx]), float(times[idx] / 3600), kind]
                    + [f"{float(c):.4f}" for c in values])
    with (folder / "table5.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
    md = ["## 表5 干基含水率 / (kg/kg)", "",
          "| 时间 / h | r=0 cm | r=0.5 cm | r=1 cm | r=1.5 cm | r=2 cm |",
          "|---|---:|---:|---:|---:|---:|"]
    for row in rows:
        label = f"{row[1]:.4f}（终点）" if row[2] == "endpoint" else f"{row[1]:g}"
        md.append("| " + " | ".join([label] + row[3:]) + " |")
    md.extend(["", f"终点原始时间为 {finish!r} s；小时数仅在表面显示时格式化。",
               "终点浓度显示为0.1500时不据此判定是否达标，严格判据使用未舍入的全域最大值。", ""])
    content = "\n".join(md)
    (folder / "table5.md").write_text(content, encoding="utf-8")
    (folder / "题目规定表格.md").write_text(content, encoding="utf-8")


def _write_diagnostics(solution, folder, times):
    logs = np.asarray(solution["logs"])
    fields = ["time_s", "time_h", "max_C", "max_r_m", "C_center", "C_surface",
              "mean_C", "loss_cumulative", "water_balance", "interval_mean_C",
              "interval_loss", "steps", "iterations", "rejections", "min_dt_s",
              "max_dt_s", "worst_scaled_residual", "damped_iterations"]
    with (folder / "solver_diagnostics.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(fields)
        for i, time in enumerate(times):
            values = [time, time / 3600, solution["max_C"][i], solution["max_r"][i],
                      solution["moisture"][i, 0], solution["moisture"][i, -1],
                      solution["mean_C"][i], solution["loss_cumulative"][i],
                      solution["water_balance"][i], *logs[i]]
            writer.writerow([float(v) if np.isfinite(v) else "" for v in values])
    with (folder / "event_trace.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "max_C", "max_r_m", "g"])
        writer.writerows(np.asarray(solution["event_trace"], dtype=float))


def _write_note(solution, folder, endpoint, workbook):
    from common.drying_result_notes import write_short_note
    write_short_note(solution, folder)


def analyze(solution, output_dir, template_path=None, result_path=None, plots=True):
    """Write Q3 deliverables; output_dir is the q3 directory, not its tables child."""
    times, radii, finish, threshold = _check_solution(solution)
    folder = Path(output_dir)
    tables = folder / "tables"
    figures = folder / "figures"
    tables.mkdir(parents=True, exist_ok=True)
    figures.mkdir(parents=True, exist_ok=True)
    root = Path(__file__).resolve().parents[1]
    template_path = Path(template_path) if template_path is not None else root / "data/templates/result3_template.xlsx"
    result_path = Path(result_path) if result_path is not None else root / "submission_results/result3.xlsx"
    workbook = export_result3(solution, template_path, result_path)
    _write_table5(solution, tables, times, radii, finish)
    _write_diagnostics(solution, tables, times)
    endpoint = {
        "finish_time_s": finish, "finish_time_h": finish / 3600,
        "threshold": threshold, "strictly_below_threshold": bool(solution["max_C"][-1] < threshold),
        "max_C": float(solution["max_C"][-1]), "max_r": float(solution["max_r"][-1]),
        "max_r_unit": "m", "g_plus": float(solution["max_C"][-1] - threshold),
        "bracket": solution["metadata"]["bracket"],
        "radii_m": radii, "temperature_C": np.asarray(solution["temperature_K"])[-1] - 273.15,
        "moisture": np.asarray(solution["moisture"])[-1],
        "sampling_note": "radii_m, temperature_C and moisture are reconstructed output samples, not native cells",
        "native_grid": {"faces_m": np.asarray(solution["faces"]),
                        "centers_m": np.asarray(solution["centers"]),
                        "volumes": np.asarray(solution["volumes"]),
                        "volume_weight_note": "Radial weights (r_outer^2-r_inner^2)/2 in m^2; common 2*pi*L omitted",
                        "temperature_K": np.asarray(solution["final_T"]),
                        "moisture": np.asarray(solution["final_C"])},
        "mean_C": float(solution["mean_C"][-1]),
        "loss_cumulative": float(solution["loss_cumulative"][-1]),
        "water_balance": float(solution["water_balance"][-1]),
        "source_hashes": solution["metadata"].get("source_hashes", {}),
        "workbook": workbook,
    }
    _write_json(folder / "endpoint.json", endpoint)
    radial_indices = _exact_indices(radii, np.arange(21) * .001, 1e-12, "endpoint radii")
    with (tables / "endpoint.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "time_h", "row_kind"] + [f"r={j * .1:g} cm" for j in range(21)])
        writer.writerow([finish, finish / 3600, "endpoint"]
                        + [f"{float(solution['moisture'][-1, j]):.4f}" for j in radial_indices])
    _write_note(solution, folder, endpoint, workbook)
    if plots:
        figure5(solution, figures)
    return {"workbook": workbook, "endpoint": str(folder / "endpoint.json"),
            "table5": str(tables / "table5.csv"),
            "figure": str(figures / "fig5_q3_drying_endpoint.pdf") if plots else None}
