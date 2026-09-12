"""Q4 exports at fixed physical radii, with separate moving-surface values."""
from pathlib import Path
import csv
import json

import bootstrap
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from q4.plotting import figure6, figure7


def _json_value(value):
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    if isinstance(value, np.ndarray):
        return _json_value(value.tolist())
    if isinstance(value, np.generic):
        return _json_value(value.item())
    if isinstance(value, float) and np.isnan(value):
        return None
    if isinstance(value, Path):
        return str(value)
    return value


def _write_json(path, value):
    Path(path).write_text(json.dumps(_json_value(value), ensure_ascii=False, indent=2,
                                    allow_nan=False) + "\n", encoding="utf-8")


def _exact_indices(available, requested, atol, name):
    available, requested = np.asarray(available), np.asarray(requested)
    index = np.minimum(np.searchsorted(available, requested), len(available) - 1)
    previous = np.maximum(index - 1, 0)
    choose = abs(available[previous] - requested) < abs(available[index] - requested)
    index[choose] = previous[choose]
    if not np.all(abs(available[index] - requested) <= atol):
        raise ValueError(f"Missing saved {name}; export interpolation is forbidden.")
    return index


def _check_solution(solution):
    times = np.asarray(solution["times"], dtype=float)
    radii = np.asarray(solution["radii"], dtype=float)
    radius = np.asarray(solution["radius_m"], dtype=float)
    xi = np.asarray(solution["xi"], dtype=float)
    if times.ndim != 1 or times[0] != 0 or not np.all(np.isfinite(times)) or not np.all(np.diff(times) > 0):
        raise ValueError("Q4 needs increasing absolute times beginning at the original initial state.")
    if radii.ndim != 1 or not np.all(np.diff(radii) > 0) or not np.allclose(radii[[0, -1]], [0, .02], atol=1e-12, rtol=0):
        raise ValueError("Fixed physical output radii must span 0 to 0.02 metres.")
    if xi.ndim != 1 or not np.all(np.diff(xi) > 0) or not np.allclose(xi[[0, -1]], [0, 1], atol=1e-12, rtol=0):
        raise ValueError("Complete reference samples must include xi=0 and xi=1.")
    if radius.shape != times.shape or not np.all(np.isfinite(radius)) or not np.all(radius > 0):
        raise ValueError("Q4 radius history must be finite and positive.")
    mask = np.asarray(solution["valid_mask"], dtype=bool)
    expected = (radii[None, :] >= 0) & (radii[None, :] <= radius[:, None])
    if mask.shape != expected.shape or not np.array_equal(mask, expected):
        raise ValueError("Valid output positions must lie inside the actual material radius.")
    for key in ("temperature_K", "moisture"):
        field = np.asarray(solution[key], dtype=float)
        if field.shape != mask.shape or not np.all(np.isfinite(field[mask])) or not np.all(np.isnan(field[~mask])):
            raise ValueError(f"{key}: material exterior must be NaN, never zero or surface-filled.")
    for key in ("temperature_ref_K", "moisture_ref"):
        field = np.asarray(solution[key], dtype=float)
        if field.shape != (len(times), len(xi)) or not np.all(np.isfinite(field)):
            raise ValueError(f"Invalid full-domain reference field {key}.")
    for key in ("max_C", "max_r", "mean_C", "loss_cumulative", "water_balance",
                "surface_temperature_K", "surface_moisture"):
        field = np.asarray(solution[key], dtype=float)
        if field.shape != times.shape or not np.all(np.isfinite(field)):
            raise ValueError(f"Invalid full-history field {key}.")
    if not np.all((np.asarray(solution["max_r"]) >= 0) & (np.asarray(solution["max_r"]) <= radius + 1e-12)):
        raise ValueError("Full-grid maximum locations must be physical positions within the moving domain.")
    if np.asarray(solution["logs"]).shape != (len(times), 9):
        raise ValueError("Q4 interval diagnostics must have nine columns.")
    finish = float(solution["metadata"]["finish_time_s"])
    threshold = float(solution["metadata"].get("threshold", .15))
    if not np.isclose(times[-1], finish, atol=1e-7, rtol=0) or not float(solution["max_C"][-1]) < threshold:
        raise ValueError("Q4 export requires the actual strictly dry terminal state.")
    bracket = solution["metadata"]["bracket"]
    if not (float(bracket["g_minus"]) >= 0 and float(bracket["g_plus"]) < 0):
        raise ValueError("The endpoint needs an undried/dried sign bracket.")
    if not np.isclose(float(bracket["t_plus"]), finish, atol=1e-7, rtol=0):
        raise ValueError("The endpoint must be the strictly dry upper bracket state.")
    return times, radii, radius, mask, finish, threshold


def export_result4(solution, template_path, destination):
    times, radii, radius, mask, finish, _ = _check_solution(solution)
    minutes = np.arange(1, int(np.floor(finish / 60)) + 1, dtype=np.int64) * 60
    rows = _exact_indices(times, minutes, 1e-7, "whole-minute times")
    columns = _exact_indices(radii, np.arange(21) * .001, 1e-12, "physical radii")
    values = np.asarray(solution["moisture"])[np.ix_(rows, columns)]
    valid = mask[np.ix_(rows, columns)]
    wb = load_workbook(template_path)
    if wb.sheetnames != ["Sheet1"]:
        raise ValueError("result4 must preserve the supplied single Sheet1 layout.")
    ws = wb["Sheet1"]
    title = ws["A1"].value or "时间\\到药材中心的距离"
    if "药材表面" not in [cell.value for cell in ws[1]]:
        raise ValueError("The original result4 template must identify a separate material surface.")
    for row in ws:
        for cell in row:
            cell.value = None
    ws["A1"] = title
    for j in range(21):
        ws.cell(1, j + 2, round(j * .1, 1)).number_format = "0.0"
    ws.cell(1, 23, "药材表面")
    for i, second in enumerate(minutes, start=2):
        ws.cell(i, 1, int(second)).number_format = "0"
        for j in range(21):
            cell = ws.cell(i, j + 2)
            cell.value = round(float(values[i - 2, j]), 4) if valid[i - 2, j] else None
            cell.number_format = "0.0000"
        ws.cell(i, 23, round(float(solution["surface_moisture"][rows[i - 2]]), 4)).number_format = "0.0000"
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 28
    ws.row_dimensions[1].height = 32
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    for j in range(2, 24):
        ws.column_dimensions[get_column_letter(j)].width = 11 if j < 23 else 14
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    wb.close()
    return {"path": str(destination), "sheet": "Sheet1", "data_rows": len(minutes),
            "fixed_physical_radial_columns": 21, "surface_column": "W", "surface_header": "药材表面",
            "first_time_s": int(minutes[0]) if len(minutes) else None,
            "last_time_s": int(minutes[-1]) if len(minutes) else None,
            "interior_fixed_values": int(valid.sum()), "exterior_blank_cells": int((~valid).sum()),
            "surface_values": len(minutes), "nonminute_endpoint_included": False}


def _concentration_text(value, valid=True):
    return f"{float(value):.4f}" if valid else ""


def _write_table6(solution, folder, times, radii, radius, mask, finish):
    positions_cm = np.array([0, .5, 1, 1.5, 2])
    columns = _exact_indices(radii, positions_cm * .01, 1e-12, "table 6 positions")
    targets = np.arange(1, int(np.floor(finish / 21600)) + 1, dtype=float) * 21600
    targets = targets[targets < finish - 1e-7]
    indices = list(_exact_indices(times, targets, 1e-7, "six-hour states")) + [len(times) - 1]
    headers = ["time_s", "time_h", "row_kind", "radius_cm"] + [f"r={r:g} cm" for r in positions_cm] + ["surface_C"]
    rows = []
    for k, index in enumerate(indices):
        kind = "endpoint" if k == len(indices) - 1 else "regular_6h"
        row = [float(times[index]), float(times[index] / 3600), kind, float(radius[index] * 100)]
        row += [_concentration_text(solution["moisture"][index, j], mask[index, j]) for j in columns]
        row.append(_concentration_text(solution["surface_moisture"][index]))
        rows.append(row)
    with (folder / "table6.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)
    lines = ["## 表6 药材干燥过程的干基含水率 / (kg/kg)", "",
             "| 时间 / h | r=0 cm | r=0.5 cm | r=1 cm | r=1.5 cm | r=2 cm | 药材表面 | R(t) / cm |",
             "|---|---:|---:|---:|---:|---:|---:|---:|"]
    for row in rows:
        time_label = f"{row[1]:.4f}（终点）" if row[2] == "endpoint" else f"{row[1]:g}"
        lines.append("| " + " | ".join([time_label] + row[4:] + [f"{row[3]:.4f}"]) + " |")
    lines += ["", "空白表示该实际径向位置位于材料域外；药材表面列始终对应当时的真实R(t)，不等于最外侧固定位置列。",
              f"精确终点原始值为 {finish!r} s；四位小数仅用于展示，严格达标用未舍入的全域最大含水率判断。", ""]
    text = "\n".join(lines)
    (folder / "table6.md").write_text(text, encoding="utf-8")
    (folder / "题目规定表格.md").write_text(text, encoding="utf-8")


def _write_diagnostics(solution, folder, times, radius, mask):
    headers = ["time_s", "time_h", "radius_m", "max_C", "max_r_m", "C_center", "C_surface",
               "mean_C", "loss_cumulative", "water_balance", "valid_fixed_positions",
               "interval_mean_C", "interval_loss", "steps", "iterations", "rejections",
               "min_dt_s", "max_dt_s", "worst_scaled_residual", "damped_iterations"]
    with (folder / "solver_diagnostics.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        for i, time in enumerate(times):
            writer.writerow([float(time), float(time / 3600), float(radius[i]),
                             float(solution["max_C"][i]), float(solution["max_r"][i]),
                             float(solution["moisture"][i, 0]), float(solution["surface_moisture"][i]),
                             float(solution["mean_C"][i]), float(solution["loss_cumulative"][i]),
                             float(solution["water_balance"][i]), int(mask[i].sum()),
                             *[float(x) for x in solution["logs"][i]]])
    with (folder / "event_trace.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "max_C", "max_r_m", "g"])
        writer.writerows(np.asarray(solution["event_trace"]))


def analyze(solution, output_dir, template_path=None, result_path=None, plots=True, comparison_solutions=None):
    times, radii, radius, mask, finish, threshold = _check_solution(solution)
    folder = Path(output_dir)
    tables, figures = folder / "tables", folder / "figures"
    tables.mkdir(parents=True, exist_ok=True)
    figures.mkdir(parents=True, exist_ok=True)
    template_path = Path(template_path) if template_path else bootstrap.ROOT / "data/templates/result4_template.xlsx"
    result_path = Path(result_path) if result_path else bootstrap.ROOT / "submission_results/result4.xlsx"
    workbook = export_result4(solution, template_path, result_path)
    _write_table6(solution, tables, times, radii, radius, mask, finish)
    _write_diagnostics(solution, tables, times, radius, mask)
    endpoint = {"finish_time_s": finish, "finish_time_h": finish / 3600,
                "threshold": threshold, "strictly_below_threshold": bool(solution["max_C"][-1] < threshold),
                "max_C": float(solution["max_C"][-1]), "max_r": float(solution["max_r"][-1]), "max_r_unit": "m",
                "g_plus": float(solution["max_C"][-1] - threshold), "bracket": solution["metadata"]["bracket"],
                "radius_m": float(radius[-1]), "radii_m": radii, "valid_mask": mask[-1],
                "temperature_C": np.asarray(solution["temperature_K"])[-1] - 273.15,
                "moisture": np.asarray(solution["moisture"])[-1],
                "surface_temperature_C": float(solution["surface_temperature_K"][-1] - 273.15),
                "surface_moisture": float(solution["surface_moisture"][-1]),
                "sampling_note": "Fixed physical output locations; exterior values are null. The actual surface is a separate reconstructed boundary value.",
                "reference_samples": {"xi": solution["xi"], "physical_radii_m": solution["reference_physical_radii"][-1],
                                      "temperature_C": np.asarray(solution["temperature_ref_K"])[-1] - 273.15,
                                      "moisture": solution["moisture_ref"][-1]},
                "native_grid": {"faces_reference_m": solution["faces"], "centers_reference_m": solution["centers"],
                                "reference_volumes": solution["volumes"],
                                "physical_faces_m": solution["final_physical_faces"],
                                "physical_centers_m": solution["final_physical_centers"],
                                "temperature_K": solution["final_T"], "moisture": solution["final_C"],
                                "volume_weight_note": "Reference radial weights in m^2; initial dry-mass normalization, not the current shrunken volume weights"},
                "mean_C": float(solution["mean_C"][-1]), "loss_cumulative": float(solution["loss_cumulative"][-1]),
                "water_balance": float(solution["water_balance"][-1]), "workbook": workbook,
                "source_hashes": solution["metadata"].get("source_hashes", {})}
    _write_json(folder / "endpoint.json", endpoint)
    columns = _exact_indices(radii, np.arange(21) * .001, 1e-12, "endpoint physical positions")
    with (tables / "endpoint.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["time_s", "time_h", "row_kind", "radius_cm"] + [f"r={j * .1:g} cm" for j in range(21)] + ["surface_C"])
        writer.writerow([finish, finish / 3600, "endpoint", float(radius[-1] * 100)]
                        + [_concentration_text(solution["moisture"][-1, j], mask[-1, j]) for j in columns]
                        + [_concentration_text(solution["surface_moisture"][-1])])
    from common.drying_result_notes import write_short_note
    write_short_note(solution, folder, comparison_solutions)
    if plots:
        figure6(solution["radius_data"], solution["radius_slopes"], figures)
        figure7(solution, figures, comparison_solutions)
    return {"workbook": workbook, "endpoint": str(folder / "endpoint.json"), "table6": str(tables / "table6.csv"),
            "figures": [str(figures / f"{stem}.pdf") for stem in ("fig6_q4_radius_mapping", "fig7_q4_shrinking_comparison")] if plots else []}
