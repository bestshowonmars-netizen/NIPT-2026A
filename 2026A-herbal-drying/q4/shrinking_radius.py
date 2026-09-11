"""Monotone PCHIP radius input in seconds/metres, with endpoint holding."""
from pathlib import Path
import numpy as np
from numba import njit
from openpyxl import load_workbook


def pchip_slopes(data):
    x, y = data[:, 0], data[:, 1]
    h = np.diff(x)
    delta = np.diff(y) / h
    d = np.zeros(len(x))
    if len(x) == 2:
        d[:] = delta[0]
        return d
    for i in range(1, len(x)-1):
        if delta[i-1] * delta[i] > 0:
            w1, w2 = 2*h[i]+h[i-1], h[i]+2*h[i-1]
            d[i] = (w1+w2)/(w1/delta[i-1]+w2/delta[i])
    for i, j, k in ((0, 0, 1), (-1, -1, -2)):
        d[i] = ((2*h[j]+h[k])*delta[j]-h[j]*delta[k])/(h[j]+h[k])
        if np.sign(d[i]) != np.sign(delta[j]):
            d[i] = 0.0
        elif np.sign(delta[j]) != np.sign(delta[k]) and abs(d[i]) > 3*abs(delta[j]):
            d[i] = 3*delta[j]
    return d


def load_radius_data(path):
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    rows = [(float(t), float(r)*0.01) for t, r, *_ in
            wb.active.iter_rows(min_row=2, values_only=True) if t is not None]
    wb.close()
    data = np.asarray(rows, dtype=np.float64)
    if (data.ndim != 2 or data.shape[1] != 2 or len(data) < 2
            or not np.all(np.isfinite(data)) or data[0, 0] != 0
            or np.any(np.diff(data[:, 0]) <= 0) or np.any(data[:, 1] <= 0)
            or np.any(np.diff(data[:, 1]) > 0)):
        raise ValueError("Radius input requires increasing seconds and positive nonincreasing metres")
    return data, pchip_slopes(data)


@njit(cache=True)
def radius_at(t, data, slopes, shrinking=True, radius0=0.02):
    if not shrinking:
        return radius0
    if t <= data[0, 0]:
        return data[0, 1]
    if t >= data[-1, 0]:
        return data[-1, 1]
    j = np.searchsorted(data[:, 0], t, side="right")-1
    # A recorded plateau is exactly constant, avoiding cancellation of basis
    # weights around an equal-radius physical sampling position.
    if data[j, 1] == data[j+1, 1]:
        return data[j, 1]
    h = data[j+1, 0]-data[j, 0]
    z = (t-data[j, 0])/h
    return ((2*z**3-3*z*z+1)*data[j, 1]
            +(z**3-2*z*z+z)*h*slopes[j]
            +(-2*z**3+3*z*z)*data[j+1, 1]
            +(z**3-z*z)*h*slopes[j+1])
